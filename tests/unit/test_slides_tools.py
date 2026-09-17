from __future__ import annotations

import asyncio
import json

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from code_ai.config.models import AppConfig
from code_ai.core.errors import ToolArgumentError
from code_ai.events.bus import AsyncEventBus
from code_ai.tools.base import TOOL_IMAGES_KEY, ToolCapability, ToolContext
from code_ai.tools.office.converters import find_libreoffice
from code_ai.tools.slides import (
    SlidesCreateTool,
    SlidesEditTool,
    SlidesFormatTool,
    SlidesInspectTool,
    SlidesRenderTool,
    fit,
    spec,
)
from code_ai.tools.slides.themes import THEMES, contrast, text_safe
from code_ai.util.paths import WorkspacePolicy

DECK = """# Data Platform 2027
## Roadmap
Engineering · September 2026
---
<!-- kind: agenda -->
# Agenda
- Context
- Proposal
---
<!-- kind: section -->
# Context
---
# Problems
- Manual **pipelines**
  - weekly breakage
- Slow queries
Note: stress the cost
---
# Before and after
<!-- kind: comparison -->
### Today
- Batch
|||
### 2027
- Streaming
---
![Architecture](photo.png)
# Architecture
- Kafka
---
# Numbers
<!-- kind: stats -->
- **-70%** latency
- **3x** users
---
# Plan
<!-- kind: timeline -->
1. **Discover** map sources
2. **Build** lakehouse
---
# Pillars
<!-- kind: cards -->
- **Reliability**: SLAs
- **Cost**: tiers
---
# Costs
| Quarter | Total |
|---|---|
| Q1 | 420 |
| Q2 | 460 |
---
# Usage
```chart
{"type":"column","categories":["Q1","Q2"],"series":{"Queries":[120,180]}}
```
---
> Reliable data drives fast decisions.
> — Engineering
---
# Ingestion
```python
for message in consumer:
    lake.write(message)
```
---
<!-- kind: closing -->
# Thank you
"""


def make_context(tmp_path) -> ToolContext:
    config = AppConfig.from_mapping({"api_mode": "ollama", "workspace": str(tmp_path)})
    return ToolContext(
        config=config,
        workspace=WorkspacePolicy.from_path(tmp_path),
        event_bus=AsyncEventBus(session_id="session"),
        cancel_event=asyncio.Event(),
    )


async def create(tmp_path, **extra):
    Image.new("RGB", (800, 500), "teal").save(tmp_path / "photo.png")
    arguments = {"output": "deck.pptx", "content": DECK, **extra}
    result = await SlidesCreateTool().execute(arguments, make_context(tmp_path))
    return result, Presentation(str(tmp_path / "deck.pptx"))


def names(slide) -> set[str]:
    return {shape.name for shape in slide.shapes}


def test_capabilities() -> None:
    assert ToolCapability.LOCAL_WRITE not in SlidesInspectTool.capabilities
    assert ToolCapability.LOCAL_WRITE in SlidesEditTool.capabilities


def test_markdown_kinds_are_inferred_and_directed() -> None:
    slides = spec.parse(DECK)
    assert [s["kind"] for s in slides] == [
        "title",
        "agenda",
        "section",
        "bullets",
        "comparison",
        "image_text",
        "stats",
        "timeline",
        "cards",
        "table",
        "chart",
        "quote",
        "code",
        "closing",
    ]
    assert slides[3]["body"] == [
        ("Manual **pipelines**", 0),
        ("weekly breakage", 1),
        ("Slow queries", 0),
    ]
    assert slides[3]["notes"] == "stress the cost"
    assert slides[4]["left"]["heading"] == "Today" and slides[4]["right"]["body"] == [
        ("Streaming", 0)
    ]
    assert slides[6]["stats"] == [("-70%", "latency"), ("3x", "users")]
    assert slides[11]["author"] == "Engineering"


def test_json_slides_and_validation() -> None:
    slides = spec.parse(
        json.dumps(
            [
                {"title": "A", "body": ["one", ["nested"]]},
                {"kind": "kpi", "title": "B", "stats": [{"value": "1", "label": "x"}]},
            ]
        )
    )
    assert slides[0]["kind"] == "bullets" and slides[0]["body"][1] == ("nested", 1)
    assert slides[1]["kind"] == "stats"
    with pytest.raises(ToolArgumentError, match="needs 'categories'"):
        spec.parse([{"kind": "chart", "chart": {"type": "column", "series": {"a": [1]}}}])
    with pytest.raises(ToolArgumentError, match="unknown kind"):
        spec.parse([{"kind": "hologram"}])


def test_fit_shrinks_long_text_and_refuses_words_wider_than_the_box() -> None:
    long = [("word " * 80, 0)]
    size, fits = fit.fit_size(long, 300, 100, "Segoe UI", start=28, minimum=10)
    assert size < 28
    assert fit.block_height([("Supercalifragilistic", 0)], 40, "Segoe UI", 24) == float("inf")


def test_text_safe_colours_reach_large_text_contrast() -> None:
    assert contrast(text_safe("14B8A6", "FFFFFF"), "FFFFFF") >= 3.0
    assert contrast(text_safe("1E293B", "0F172A"), "0F172A") >= 3.0


@pytest.mark.parametrize("theme", sorted(THEMES))
async def test_every_theme_builds_every_kind_without_warnings(tmp_path, theme) -> None:
    result, deck = await create(tmp_path, theme=theme, footer="Platform")
    assert result["slides"] == 14 and "warnings" not in result
    assert deck.core_properties.category == f"code-ai-theme:{theme}"
    chart_slide = deck.slides[10]
    assert any(shape.has_chart for shape in chart_slide.shapes)
    table = next(shape for shape in deck.slides[9].shapes if shape.has_table)
    assert table.table.cell(1, 0).text == "Q1"
    assert deck.slides[3].notes_slide.notes_text_frame.text == "stress the cost"
    assert "Slide Number" in names(deck.slides[3]) and "Footer" in names(deck.slides[3])
    picture = next(s for s in deck.slides[5].shapes if s.shape_type == 13)
    assert picture._element.nvPicPr.cNvPr.get("descr") == "Architecture"


async def test_our_decks_pass_our_own_inspection(tmp_path) -> None:
    await create(tmp_path)
    report = await SlidesInspectTool().execute({"path": "deck.pptx"}, make_context(tmp_path))
    rules = {issue["rule"] for issue in report["issues"]}
    assert rules <= {"missing-title"}  # the quote slide has no title by design
    detail = report["slide_details"][10]
    chart = next(shape for shape in detail["shapes"] if shape["kind"] == "chart")
    assert chart["series"] == {"Queries": [120.0, 180.0]}


async def test_inspect_flags_a_messy_deck(tmp_path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Results"
    box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(3), Inches(0.6))
    box.text_frame.word_wrap = True
    box.text_frame.text = (
        "A very long sentence that will never fit in such a small box at this size " * 2
    )
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0xEE, 0xEE, 0xEE)
    tiny = deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Inches(9.5), Inches(1), Inches(2), Inches(1)
    )
    tiny.text_frame.text = "tiny"
    tiny.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    deck.save(tmp_path / "messy.pptx")

    report = await SlidesInspectTool().execute({"path": "messy.pptx"}, make_context(tmp_path))
    rules = {issue["rule"] for issue in report["issues"]}
    assert {
        "text-overflow",
        "low-contrast",
        "empty-placeholder",
        "small-text",
        "off-slide",
        "missing-title",
    } <= rules

    result = await SlidesFormatTool().execute(
        {"path": "messy.pptx", "output": "clean.pptx"}, make_context(tmp_path)
    )
    assert result["changes"]["empty placeholders removed"] == 1
    assert "low-contrast text recoloured" in result["changes"]
    after = await SlidesInspectTool().execute({"path": "clean.pptx"}, make_context(tmp_path))
    left = {issue["rule"] for issue in after["issues"]}
    assert not left & {"low-contrast", "empty-placeholder", "small-text", "text-overflow"}


async def test_edit_operations(tmp_path) -> None:
    await create(tmp_path)
    operations = [
        {"op": "duplicate", "slide": 11},
        {
            "op": "update_chart",
            "slide": 12,
            "categories": ["Q1", "Q2"],
            "series": {"Queries": [1, 2]},
        },
        {"op": "duplicate", "slide": 6},
        {"op": "replace_text", "find": "Kafka", "replace": "Redpanda"},
        {
            "op": "add",
            "at": 2,
            "slide": "# Risks\n<!-- kind: cards -->\n- **Time**: vendors\n- **Cost**: FX",
        },
        {"op": "set_text", "slide": 5, "shape": "Body", "text": "- First\n- Second\n  - detail"},
        {"op": "notes", "slide": 1, "text": "Open with the incident."},
        {
            "op": "add_text",
            "slide": 1,
            "text": "Draft",
            "x": 10,
            "y": 0.2,
            "width": 2,
            "height": 0.5,
        },
        {"op": "move", "slide": "last", "to": 1},
        {"op": "delete", "slides": "1"},
        {"op": "hide", "slides": "3"},
    ]
    result = await SlidesEditTool().execute(
        {"path": "deck.pptx", "operations": json.dumps(operations)}, make_context(tmp_path)
    )
    assert result["slides"] == 16
    deck = Presentation(str(tmp_path / "deck.pptx"))
    titles = [
        next((s.text_frame.text for s in slide.shapes if s.name == "Title"), "")
        for slide in deck.slides
    ]
    assert titles[1] == "Risks"
    charts = [
        list(shape.chart.plots[0].series[0].values)
        for slide in deck.slides
        for shape in slide.shapes
        if shape.has_chart
    ]
    assert sorted(charts) == [[1.0, 2.0], [120.0, 180.0]]  # the copy is independent
    texts = " ".join(
        s.text_frame.text for slide in deck.slides for s in slide.shapes if s.has_text_frame
    )
    assert "Redpanda" in texts and "Kafka" not in texts
    body = next(s for s in deck.slides[4].shapes if s.name == "Body")
    assert [p.text for p in body.text_frame.paragraphs] == ["First", "Second", "detail"]
    assert deck.slides[2]._element.get("show") == "0"
    pictures = [s for slide in deck.slides for s in slide.shapes if s.shape_type == 13]
    assert len(pictures) == 2
    # The duplicated deck must still open and save.
    deck.save(tmp_path / "again.pptx")


async def test_bad_edits_leave_the_deck_alone(tmp_path) -> None:
    await create(tmp_path)
    before = (tmp_path / "deck.pptx").read_bytes()
    context = make_context(tmp_path)
    with pytest.raises(ToolArgumentError, match="unknown op"):
        await SlidesEditTool().execute(
            {"path": "deck.pptx", "operations": '[{"op":"explode"}]'}, context
        )
    with pytest.raises(ToolArgumentError, match="no shape"):
        await SlidesEditTool().execute(
            {
                "path": "deck.pptx",
                "operations": '[{"op":"set_text","slide":2,"shape":"Nope","text":"x"}]',
            },
            context,
        )
    assert (tmp_path / "deck.pptx").read_bytes() == before


async def test_template_masters_are_reused(tmp_path) -> None:
    template = Presentation()
    template.slide_width, template.slide_height = Inches(10), Inches(7.5)
    template.slides.add_slide(template.slide_layouts[0])
    template.save(tmp_path / "brand.pptx")
    Image.new("RGB", (800, 500), "teal").save(tmp_path / "photo.png")
    await SlidesCreateTool().execute(
        {"output": "deck.pptx", "content": "# Hello\n- world", "template": "brand.pptx"},
        make_context(tmp_path),
    )
    deck = Presentation(str(tmp_path / "deck.pptx"))
    assert len(deck.slides) == 1 and deck.slide_width == Inches(10)


@pytest.mark.skipif(find_libreoffice() is None, reason="LibreOffice not installed")
async def test_render_returns_images_and_exports_pdf(tmp_path) -> None:
    await create(tmp_path)
    result = await SlidesRenderTool().execute(
        {"path": "deck.pptx", "slides": "1-2", "export_pdf": "deck.pdf", "max_side": 400},
        make_context(tmp_path),
    )
    assert len(result[TOOL_IMAGES_KEY]) == 2
    assert (tmp_path / "deck.pdf").stat().st_size > 10_000
