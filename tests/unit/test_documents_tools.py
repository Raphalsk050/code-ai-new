from __future__ import annotations

import asyncio
import json

import docx
import pytest
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

from code_ai.config.models import AppConfig
from code_ai.core.errors import ToolArgumentError
from code_ai.events.bus import AsyncEventBus
from code_ai.tools.base import TOOL_IMAGES_KEY, ToolCapability, ToolContext
from code_ai.tools.documents import (
    DocumentConvertTool,
    DocumentCreateTool,
    DocumentEditTool,
    DocumentFormatTool,
    DocumentInspectTool,
)
from code_ai.tools.documents.ooxml import list_info
from code_ai.tools.documents.presets import PRESETS
from code_ai.tools.office.converters import find_libreoffice
from code_ai.util.paths import WorkspacePolicy

MARKDOWN = """# Report Title

## Context

The system uses a **monolith** and *must* migrate. See [docs](https://example.com) and `code`.

1. First
2. Second
   - nested a
   - nested b
3. Third

- Loose bullet

> A quoted passage.

Table: Schedule

| Phase | Start | End |
|---|:-:|--:|
| Discovery | jan | feb |
| Build | mar | jun |

![Target architecture](fig.png){width=8cm}

```python
print("hi")
```

\\newpage

## Risks

1. Restarted list
"""


def make_context(tmp_path) -> ToolContext:
    config = AppConfig.from_mapping({"api_mode": "ollama", "workspace": str(tmp_path)})
    return ToolContext(
        config=config,
        workspace=WorkspacePolicy.from_path(tmp_path),
        event_bus=AsyncEventBus(session_id="session"),
        cancel_event=asyncio.Event(),
    )


async def create(tmp_path, **extra) -> docx.Document:
    Image.new("RGB", (400, 200), "navy").save(tmp_path / "fig.png")
    arguments = {"output": "out.docx", "content": MARKDOWN, **extra}
    await DocumentCreateTool().execute(arguments, make_context(tmp_path))
    return docx.Document(str(tmp_path / "out.docx"))


def styles(document) -> list[str]:
    return [p.style.name for p in document.paragraphs]


def test_capabilities() -> None:
    assert ToolCapability.LOCAL_WRITE not in DocumentInspectTool.capabilities
    assert ToolCapability.LOCAL_WRITE in DocumentEditTool.capabilities


async def test_markdown_becomes_real_word_structure(tmp_path) -> None:
    document = await create(tmp_path)
    names = styles(document)
    assert names[0] == "Title"
    assert names.count("Heading 1") == 2  # '##' is the first level under a lone '#'
    assert "Quote" in names and "Code Block" in names and names.count("Caption") == 2

    items = [(p.text, list_info(p)) for p in document.paragraphs if list_info(p)]
    assert items[0] == ("First", ("number", 0))
    assert items[2] == ("nested a", ("bullet", 1))
    first_num = document.paragraphs[names.index("List Paragraph")]._p.pPr.numPr.numId.val
    restarted = [p for p in document.paragraphs if p.text == "Restarted list"][0]
    assert restarted._p.pPr.numPr.numId.val != first_num

    context = document.paragraphs[2]
    assert context._p.find(qn("w:hyperlink")) is not None
    assert any(run.bold for run in context.runs)
    assert len(document.inline_shapes) == 1
    assert document.inline_shapes[0]._inline.docPr.get("descr") == "Target architecture"
    table = document.tables[0]
    assert table.cell(1, 0).text == "Discovery"
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert any(br.get(qn("w:type")) == "page" for br in document.element.body.iter(qn("w:br")))


@pytest.mark.parametrize("preset", sorted(PRESETS))
async def test_every_preset_builds(tmp_path, preset) -> None:
    document = await create(tmp_path, preset=preset, toc=True)
    normal = document.styles["Normal"]
    assert normal.font.name == PRESETS[preset].body_font
    assert normal.font.size == Pt(PRESETS[preset].body_size)
    assert any("TOC" in (t.text or "") for t in document.element.body.iter(qn("w:instrText")))


async def test_abnt_cover_margins_and_numbering(tmp_path) -> None:
    cover = {
        "institution": "Universidade",
        "author": "Ana",
        "title": "Estudo",
        "city": "Manaus",
        "year": "2026",
    }
    document = await create(tmp_path, preset="abnt", cover=json.dumps(cover))
    section = document.sections[0]
    assert round(section.left_margin.cm, 1) == 3.0 and round(section.right_margin.cm, 1) == 2.0
    assert section.different_first_page_header_footer
    assert document.paragraphs[0].text == "UNIVERSIDADE"
    heading = document.styles["Heading 1"]
    assert heading.element.pPr.find(qn("w:numPr")) is not None
    assert heading.font.all_caps
    assert document.core_properties.title == "Estudo"
    captions = [p.text for p in document.paragraphs if p.style.name == "Caption"]
    assert captions[0].startswith("Tabela")


async def test_header_footer_fields(tmp_path) -> None:
    document = await create(tmp_path, footer="Page {page} of {total}", preset="corporate")
    footer = document.sections[0].footer._element
    codes = " ".join(t.text for t in footer.iter(qn("w:instrText")))
    assert "PAGE" in codes and "NUMPAGES" in codes


async def test_inspect_lists_blocks_and_issues(tmp_path) -> None:
    document = docx.Document()
    document.add_paragraph("INTRODUCTION").runs[0].bold = True
    wrong = document.add_paragraph("Body in the wrong font that goes on for a while.")
    wrong.runs[0].font.name = "Comic Sans MS"
    wrong.runs[0].font.size = Pt(15)
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("1. typed item")
    document.add_paragraph("2. typed item")
    document.add_paragraph("More body text in the default font for comparison here.")
    document.save(tmp_path / "messy.docx")

    report = await DocumentInspectTool().execute({"path": "messy.docx"}, make_context(tmp_path))
    rules = {issue["rule"] for issue in report["issues"]}
    assert {"fake-heading", "manual-numbering", "direct-formatting", "empty-paragraphs"} <= rules
    assert report["blocks"][0].startswith("p0 [Normal]")


async def test_format_fixes_the_mess(tmp_path) -> None:
    document = docx.Document()
    document.add_paragraph("OVERVIEW").runs[0].bold = True
    wrong = document.add_paragraph("Body text in a stray font.")
    wrong.runs[0].font.name = "Comic Sans MS"
    document.add_paragraph("")
    document.add_paragraph("- typed bullet")
    document.add_paragraph("- another")
    document.add_paragraph("Figure 1 - a caption")
    document.add_paragraph("Closing text.")
    document.save(tmp_path / "messy.docx")

    result = await DocumentFormatTool().execute(
        {"path": "messy.docx", "preset": "corporate"}, make_context(tmp_path)
    )
    changes = result["changes"]
    assert changes["fake headings turned into real headings"] == 1
    assert changes["typed list items turned into real lists"] == 2
    assert changes["empty spacer paragraphs removed"] == 1
    formatted = docx.Document(str(tmp_path / "messy.docx"))
    assert formatted.paragraphs[0].style.name == "Heading 1"
    assert formatted.paragraphs[1].runs[0].font.name is None
    assert list_info(formatted.paragraphs[2]) == ("bullet", 0)
    assert formatted.paragraphs[2].text == "typed bullet"
    assert formatted.paragraphs[4].style.name == "Caption"


async def test_edit_operations(tmp_path) -> None:
    await create(tmp_path, preset="default")
    before = docx.Document(str(tmp_path / "out.docx"))
    assert "monolith" in before.paragraphs[2].text
    operations = [
        {"op": "replace", "find": "uses a monolith", "replace": "uses services"},
        {"op": "format_text", "find": "Third", "bold": True, "highlight": "yellow"},
        {"op": "insert", "after_heading": "Context", "markdown": "Inserted *paragraph*."},
        {"op": "table_add_row", "table": 0, "values": ["Test", "jul", "aug"]},
        {"op": "table_cell", "table": 0, "row": 1, "column": 0, "text": "Research"},
        {"op": "comment", "paragraph": 2, "find": "system", "text": "check"},
        {"op": "delete", "section": "Risks"},
        {"op": "properties", "title": "v2", "author": "Ana"},
        {"op": "header_footer", "footer": "{page}", "position": "footer-right"},
    ]
    result = await DocumentEditTool().execute(
        {"path": "out.docx", "operations": json.dumps(operations)}, make_context(tmp_path)
    )
    assert len(result["applied"]) == len(operations)
    document = docx.Document(str(tmp_path / "out.docx"))
    texts = [p.text for p in document.paragraphs]
    assert "The system uses services and must migrate." in texts[2].replace("  ", " ")
    assert "Inserted paragraph." in texts
    # Inserted at the end of the section, before the page break that precedes Risks.
    page_break = next(
        i
        for i, p in enumerate(document.paragraphs)
        if any(br.get(qn("w:type")) == "page" for br in p._p.iter(qn("w:br")))
    )
    assert texts.index("Inserted paragraph.") < page_break
    assert not any("Risks" in text for text in texts)
    third = [p for p in document.paragraphs if "Third" in p.text][0]
    assert any(run.bold and run.text == "Third" for run in third.runs)
    assert document.tables[0].cell(1, 0).text == "Research"
    assert len(document.tables[0].rows) == 4
    assert document.core_properties.title == "v2"
    assert len(list(document.comments)) == 1


async def test_replace_spans_runs_and_keeps_first_run_format(tmp_path) -> None:
    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Hel").bold = True
    paragraph.add_run("lo wor")
    paragraph.add_run("ld!").italic = True
    document.save(tmp_path / "runs.docx")
    await DocumentEditTool().execute(
        {
            "path": "runs.docx",
            "operations": [{"op": "replace", "find": "Hello world", "replace": "Hi"}],
        },
        make_context(tmp_path),
    )
    runs = docx.Document(str(tmp_path / "runs.docx")).paragraphs[0].runs
    assert "".join(run.text for run in runs) == "Hi!"
    assert runs[0].bold


async def test_tracked_changes_accept(tmp_path) -> None:
    document = docx.Document()
    paragraph = document.add_paragraph("kept ")
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    from lxml import etree

    inserted = etree.fromstring(
        f'<w:ins xmlns:w="{w}" w:id="1" w:author="a"><w:r><w:t>added</w:t></w:r></w:ins>'
    )
    deleted = etree.fromstring(
        f'<w:del xmlns:w="{w}" w:id="2" w:author="a"><w:r><w:delText>gone</w:delText></w:r></w:del>'
    )
    paragraph._p.append(inserted)
    paragraph._p.append(deleted)
    document.save(tmp_path / "tracked.docx")
    await DocumentEditTool().execute(
        {"path": "tracked.docx", "operations": [{"op": "accept_changes"}]}, make_context(tmp_path)
    )
    accepted = docx.Document(str(tmp_path / "tracked.docx"))
    assert accepted.paragraphs[0].text == "kept added"


async def test_bad_operations_do_not_touch_the_file(tmp_path) -> None:
    await create(tmp_path)
    before = (tmp_path / "out.docx").read_bytes()
    context = make_context(tmp_path)
    with pytest.raises(ToolArgumentError, match="unknown op"):
        await DocumentEditTool().execute(
            {"path": "out.docx", "operations": '[{"op":"zap"}]'}, context
        )
    with pytest.raises(ToolArgumentError, match="does not exist"):
        await DocumentEditTool().execute(
            {"path": "out.docx", "operations": '[{"op":"delete","paragraphs":"999"}]'}, context
        )
    with pytest.raises(ToolArgumentError, match="no heading"):
        await DocumentEditTool().execute(
            {"path": "out.docx", "operations": '[{"op":"delete","section":"Nope"}]'}, context
        )
    assert (tmp_path / "out.docx").read_bytes() == before


async def test_docx_markdown_round_trip(tmp_path) -> None:
    await create(tmp_path)
    context = make_context(tmp_path)
    result = await DocumentConvertTool().execute(
        {"source": "out.docx", "output": "back.md"}, context
    )
    markdown = (tmp_path / "back.md").read_text(encoding="utf-8")
    assert markdown.startswith("# Report Title")
    assert "## Context" in markdown and "**monolith**" in markdown
    assert "[docs](https://example.com)" in markdown
    assert "   - nested a" in markdown and "| Discovery | jan | feb |" in markdown
    assert result["images"] == 1 and (tmp_path / "back_media" / "image1.png").exists()

    (tmp_path / "again.md").write_text(markdown, encoding="utf-8")
    await DocumentConvertTool().execute({"source": "again.md", "output": "again.docx"}, context)
    assert docx.Document(str(tmp_path / "again.docx")).tables


async def test_missing_images_are_reported_not_fatal(tmp_path) -> None:
    result = await DocumentCreateTool().execute(
        {"output": "a.docx", "content": "![x](nope.png)\n\n![y](https://example.com/a.png)"},
        make_context(tmp_path),
    )
    assert len(result["notes"]) == 2


@pytest.mark.skipif(find_libreoffice() is None, reason="LibreOffice not installed")
async def test_render_and_pdf_conversion(tmp_path) -> None:
    Image.new("RGB", (400, 200), "navy").save(tmp_path / "fig.png")
    result = await DocumentCreateTool().execute(
        {"output": "out.docx", "content": MARKDOWN, "render_pages": "1"}, make_context(tmp_path)
    )
    assert len(result[TOOL_IMAGES_KEY]) == 1 and result["total_pages"] >= 2
    converted = await DocumentConvertTool().execute(
        {"source": "out.docx", "output": "out.pdf"}, make_context(tmp_path)
    )
    assert converted["via"] == "libreoffice" and (tmp_path / "out.pdf").stat().st_size > 1000
