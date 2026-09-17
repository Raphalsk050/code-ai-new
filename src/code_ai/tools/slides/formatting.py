"""Bringing an existing deck into line: fonts, colours, title placement, sizes, overflow."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from typing import Any

from pptx.util import Emu, Inches, Pt

from code_ai.tools.slides import fit
from code_ai.tools.slides.inspect import (
    background_hex,
    estimated_overflow,
    is_title,
    shape_kind,
)
from code_ai.tools.slides.layouts import Grid, Painter
from code_ai.tools.slides.themes import Theme, contrast, readable_on, rgb

_CHROME = {"footer", "slide number"}


@dataclass
class SlideFormatOptions:
    apply_fonts: bool = True
    apply_background: bool = False
    fix_contrast: bool = True
    min_font_size: float = 12
    fix_overflow: bool = True
    unify_titles: bool = True
    remove_empty_placeholders: bool = True
    fill_alt_text: bool = True
    slide_numbers: bool = False


def format_deck(deck, theme: Theme, options: SlideFormatOptions) -> dict[str, Any]:
    changes: Counter[str] = Counter()
    title_geometry = _common_title_geometry(deck) if options.unify_titles else None
    for number, slide in enumerate(deck.slides, start=1):
        if options.apply_background:
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = rgb(theme.background)
            changes["backgrounds set"] += 1
        background = background_hex(slide)
        for shape in list(slide.shapes):
            name = shape.name.strip().lower()
            if (
                options.remove_empty_placeholders
                and shape.is_placeholder
                and shape.has_text_frame
                and not shape.text_frame.text.strip()
            ):
                if shape_kind(shape) == "placeholder":
                    shape._element.getparent().remove(shape._element)
                    changes["empty placeholders removed"] += 1
                    continue
            if options.fill_alt_text and shape_kind(shape) == "picture":
                element = shape._element.nvPicPr.cNvPr
                if not (element.get("descr") or "").strip():
                    element.set("descr", shape.name)
                    changes["pictures given alt text"] += 1
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        _format_frame(
                            cell.text_frame,
                            theme,
                            options,
                            changes,
                            title=False,
                            fill=_cell_fill(cell) or background,
                        )
                continue
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            title = is_title(shape)
            if title and title_geometry is not None:
                left, top, width, height = title_geometry
                if (
                    (shape.left, shape.top) != (left, top)
                    and abs(shape.left - left) < Inches(2.5)
                    and abs(shape.top - top) < Inches(1.5)
                ):
                    shape.left, shape.top, shape.width, shape.height = left, top, width, height
                    changes["titles aligned to the common position"] += 1
            fill = _shape_fill(shape) or background
            _format_frame(
                shape.text_frame,
                theme,
                options,
                changes,
                title=title,
                fill=fill,
                chrome=name in _CHROME,
            )
            if options.fix_overflow and estimated_overflow(shape) > 6:
                if _shrink(shape, options.min_font_size):
                    changes["overflowing text boxes shrunk to fit"] += 1
                elif _grow(shape, deck.slide_height):
                    changes["overflowing text boxes shrunk and extended downwards"] += 1
                else:
                    changes["text boxes still overflowing at the minimum size (split them)"] += 1
        if options.slide_numbers and not any(s.name == "Slide Number" for s in slide.shapes):
            grid = Grid(deck.slide_width / 914400, deck.slide_height / 914400)
            painter = Painter(slide, theme, grid, lambda s: s, number, show_number=True)
            painter.footer_bar(readable_on(background, theme.muted))
            changes["slide numbers added"] += 1
    return {key: value for key, value in changes.items() if value}


def _format_frame(
    frame,
    theme: Theme,
    options: SlideFormatOptions,
    changes: Counter,
    *,
    title: bool,
    fill: str,
    chrome: bool = False,
) -> None:
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            if options.apply_fonts:
                code = (run.font.name or "").lower() in {
                    theme.code_font.lower(),
                    "consolas",
                    "courier new",
                }
                wanted = theme.code_font if code else theme.title_font if title else theme.body_font
                if run.font.name != wanted:
                    run.font.name = wanted
                    changes["runs set to the theme fonts"] += 1
            if (
                not chrome
                and run.font.size is not None
                and run.font.size.pt < options.min_font_size
            ):
                run.font.size = Pt(options.min_font_size)
                changes[f"text raised to {options.min_font_size:g}pt"] += 1
            if options.fix_contrast:
                color = None
                try:
                    if run.font.color.type is not None:
                        color = str(run.font.color.rgb)
                except (AttributeError, TypeError, ValueError):
                    color = None
                if color and contrast(color, fill) < 3.0:
                    run.font.color.rgb = rgb(
                        readable_on(fill, theme.text if not title else theme.text)
                    )
                    changes["low-contrast text recoloured"] += 1


def _shape_fill(shape) -> str | None:
    try:
        if shape.fill.type == 1:
            return str(shape.fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError):
        return None
    return None


def _cell_fill(cell) -> str | None:
    try:
        if cell.fill.type == 1:
            return str(cell.fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError):
        return None
    return None


def _shrink(shape, minimum: float) -> bool:
    runs = [run for p in shape.text_frame.paragraphs for run in p.runs if run.text.strip()]
    if not runs:
        return True
    for _ in range(40):
        if estimated_overflow(shape) <= 6:
            return True
        sizes = [run.font.size.pt if run.font.size else 18.0 for run in runs]
        if max(sizes) <= minimum:
            return False
        for run, size in zip(runs, sizes, strict=True):
            run.font.size = Pt(max(minimum, size - (1 if size > 14 else 0.5)))
    return estimated_overflow(shape) <= 6


def _grow(shape, slide_height) -> bool:
    """Extend the box towards the bottom margin, just as far as the text needs."""

    limit = slide_height - Inches(0.4) - shape.top
    needed = shape.height + Pt(estimated_overflow(shape) + 4)
    if needed > limit:
        shape.height = max(shape.height, limit)
        return False
    shape.height = needed
    return estimated_overflow(shape) <= 6


def _common_title_geometry(deck) -> tuple[Emu, Emu, Emu, Emu] | None:
    boxes: Counter[tuple] = Counter()
    for slide in deck.slides:
        for shape in slide.shapes:
            if is_title(shape) and shape.has_text_frame and shape.top is not None:
                boxes[(shape.left, shape.top, shape.width, shape.height)] += 1
    if not boxes:
        return None
    geometry, count = boxes.most_common(1)[0]
    return geometry if count >= 2 else None


_unused = fit
