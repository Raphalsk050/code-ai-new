"""Opening, building and saving decks."""

from __future__ import annotations

import os
from collections.abc import Callable
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches

from code_ai.core.errors import ToolArgumentError, ToolExecutionError
from code_ai.tools.slides import layouts
from code_ai.tools.slides.themes import Theme

EMU_PER_INCH = 914400
SIZES = {"16:9": (13.333, 7.5), "4:3": (10.0, 7.5), "16:10": (12.8, 8.0)}


def open_deck(path: Path):
    try:
        return Presentation(str(path))
    except Exception as exc:  # noqa: BLE001 - python-pptx raises several types on bad files
        raise ToolExecutionError(f"Could not open {path.name} as a presentation: {exc}") from exc


def new_deck(aspect: str = "16:9", template: Path | None = None):
    if template is not None:
        deck = open_deck(template)
        # Keep masters, layouts and theme; drop the example slides.
        slide_ids = deck.slides._sldIdLst
        for slide_id in list(slide_ids):
            deck.part.drop_rel(slide_id.rId)
            slide_ids.remove(slide_id)
        return deck
    if aspect not in SIZES:
        raise ToolArgumentError(f"aspect is one of {', '.join(SIZES)}.")
    deck = Presentation()
    width, height = SIZES[aspect]
    deck.slide_width = Inches(width)
    deck.slide_height = Inches(height)
    return deck


def grid_for(deck) -> layouts.Grid:
    return layouts.Grid(deck.slide_width / EMU_PER_INCH, deck.slide_height / EMU_PER_INCH)


def blank_layout(deck):
    """The layout with the fewest placeholders - "Blank" in every stock template."""

    for layout in deck.slide_layouts:
        if layout.name.strip().lower() == "blank":
            return layout
    return min(deck.slide_layouts, key=lambda layout: len(layout.placeholders))


def add_slide(
    deck,
    spec: dict[str, Any],
    theme: Theme,
    *,
    resolve_image: Callable[[str], Path],
    number: int,
    footer: str = "",
    show_number: bool = True,
    position: int | None = None,
) -> list[str]:
    slide = deck.slides.add_slide(blank_layout(deck))
    for placeholder in list(slide.placeholders):
        placeholder._element.getparent().remove(placeholder._element)
    painter = layouts.Painter(
        slide=slide,
        theme=theme,
        grid=grid_for(deck),
        resolve_image=resolve_image,
        number=number,
        footer=footer,
        show_number=show_number and spec["kind"] not in {"title", "closing"},
    )
    layouts.draw(painter, spec)
    if position is not None:
        move_slide(deck, len(deck.slides) - 1, position)
    return painter.warnings


def move_slide(deck, old: int, new: int) -> None:
    slide_ids = deck.slides._sldIdLst
    items = list(slide_ids)
    element = items[old]
    slide_ids.remove(element)
    slide_ids.insert(max(0, min(new, len(items) - 1)), element)


def save_deck(deck, target: Path) -> None:
    staging = target.with_name(f".{target.stem}.{os.getpid()}.tmp{target.suffix}")
    try:
        deck.save(str(staging))
        os.replace(staging, target)
    finally:
        staging.unlink(missing_ok=True)
