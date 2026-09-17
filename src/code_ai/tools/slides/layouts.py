"""Drawing each slide kind on a shared grid, in the deck's theme.

Every layout places its own shapes on a blank slide rather than relying on a
template's placeholders, so a deck looks the same in PowerPoint, LibreOffice
and Google Slides.
"""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from lxml import etree
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from code_ai.core.errors import ToolArgumentError
from code_ai.tools.slides import fit
from code_ai.tools.slides.spec import plain, runs
from code_ai.tools.slides.themes import Theme, readable_on, rgb, text_safe

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}

CHARTS = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "area": XL_CHART_TYPE.AREA,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
}


@dataclass
class Grid:
    width: float
    height: float
    margin: float = 0.75

    @property
    def content_top(self) -> float:
        return 1.85

    @property
    def content_bottom(self) -> float:
        return self.height - 0.8

    @property
    def content_width(self) -> float:
        return self.width - 2 * self.margin

    @property
    def content_height(self) -> float:
        return self.content_bottom - self.content_top


@dataclass
class Painter:
    slide: Any
    theme: Theme
    grid: Grid
    resolve_image: Callable[[str], Path]
    number: int
    warnings: list[str] = field(default_factory=list)
    footer: str = ""
    show_number: bool = True

    # -- primitives -----------------------------------------------------------------------

    def background(self, color: str) -> None:
        fill = self.slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(color)

    def rect(self, x, y, w, h, color: str, *, rounded: bool = False, name: str = "", line=None):
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(w),
            Inches(h),
        )
        if rounded:
            shape.adjustments[0] = min(0.5, 0.12 / max(0.2, min(w, h)))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        if line:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        if name:
            shape.name = name
        return shape

    def oval(self, x, y, d, color: str):
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def text(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        paragraphs: list[tuple[str, int]] | str,
        *,
        size: float,
        minimum: float | None = None,
        font: str | None = None,
        color: str | None = None,
        bold: bool = False,
        italic: bool = False,
        align: str = "left",
        anchor: str = "top",
        bullets: bool = False,
        gap: float = 0.45,
        name: str = "",
        what: str = "text",
    ):
        if isinstance(paragraphs, str):
            paragraphs = [(line, 0) for line in paragraphs.split("\n")] or [("", 0)]
        font = font or self.theme.body_font
        color = color or self.theme.text
        indent = 0.32 * 72 if bullets else 0.0
        measured = [(plain(t), level) for t, level in paragraphs]
        final = size
        if minimum is not None:
            final, fits = fit.fit_size(
                measured,
                w * 72,
                h * 72,
                font,
                start=size,
                minimum=minimum,
                bold=bold,
                indent=indent,
                gap=gap,
            )
            if not fits:
                self.warnings.append(
                    f"slide {self.number}: {what} is still too long at {minimum:g}pt - "
                    "shorten it or split the slide"
                )
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            box.name = name
        frame = box.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = ANCHOR[anchor]
        for index, (content, level) in enumerate(paragraphs):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            level_size = final * (0.88 ** min(level, 3))
            paragraph.alignment = ALIGN[align]
            paragraph.line_spacing = 1.05
            if index:
                paragraph.space_before = Pt(level_size * gap)
            if bullets:
                self._bullet(paragraph, level, level_size)
            for segment, seg_bold, seg_italic, code in runs(content):
                run = paragraph.add_run()
                run.text = segment
                run.font.size = Pt(level_size)
                run.font.name = self.theme.code_font if code else font
                run.font.bold = bold or seg_bold
                run.font.italic = italic or seg_italic
                run.font.color.rgb = rgb(self.theme.accent if code and not bold else color)
        return box

    def _bullet(self, paragraph, level: int, size: float) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        step = Inches(0.32)
        p_pr.set("marL", str(int(step * (level + 1))))
        p_pr.set("indent", str(-int(step * 0.9)))
        for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buNone"):
            for old in p_pr.findall(qn(tag)):
                p_pr.remove(old)
        color = etree.SubElement(p_pr, qn("a:buClr"))
        srgb = etree.SubElement(color, qn("a:srgbClr"))
        srgb.set("val", self.theme.accent if level == 0 else self.theme.muted)
        etree.SubElement(p_pr, qn("a:buSzPct")).set("val", "100000")
        font = etree.SubElement(p_pr, qn("a:buFont"))
        font.set("typeface", "Arial")
        etree.SubElement(p_pr, qn("a:buChar")).set("char", "•" if level == 0 else "–")

    def picture(
        self,
        src: str,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        cover: bool = False,
        alt: str = "",
    ):
        try:
            path = self.resolve_image(src)
        except Exception as exc:  # noqa: BLE001 - reported on the slide instead
            self.warnings.append(f"slide {self.number}: image {src!r} not found ({exc})")
            box = self.rect(x, y, w, h, self.theme.surface, rounded=True)
            self.text(
                x,
                y,
                w,
                h,
                f"[image: {src}]",
                size=14,
                color=self.theme.muted,
                align="center",
                anchor="middle",
            )
            return box
        from PIL import Image

        try:
            with Image.open(path) as image:
                px_w, px_h = image.size
        except OSError as exc:
            raise ToolArgumentError(
                f"slide {self.number}: cannot read image {src}: {exc}"
            ) from None
        ratio = px_w / px_h
        if cover:
            picture = self.slide.shapes.add_picture(
                str(path), Inches(x), Inches(y), Inches(w), Inches(h)
            )
            box_ratio = w / h
            if ratio > box_ratio:
                crop = (1 - box_ratio / ratio) / 2
                picture.crop_left = picture.crop_right = crop
            else:
                crop = (1 - ratio / box_ratio) / 2
                picture.crop_top = picture.crop_bottom = crop
        else:
            fw, fh = (w, w / ratio) if ratio > w / h else (h * ratio, h)
            picture = self.slide.shapes.add_picture(
                str(path),
                Inches(x + (w - fw) / 2),
                Inches(y + (h - fh) / 2),
                Inches(fw),
                Inches(fh),
            )
        picture._element.nvPicPr.cNvPr.set("descr", alt or Path(src).stem)
        return picture

    # -- chrome ---------------------------------------------------------------------------

    def chrome(self, title: str, subtitle: str = "") -> float:
        """Background, title, accent and footer; returns where content may start."""

        theme, grid = self.theme, self.grid
        self.background(theme.background)
        title_h = 0.95
        self.text(
            grid.margin,
            0.45,
            grid.content_width,
            title_h,
            title or "",
            size=theme.title_size,
            minimum=22,
            font=theme.title_font,
            bold=theme.title_bold,
            anchor="bottom",
            name="Title",
            what="title",
        )
        top = 0.45 + title_h
        if theme.accent_bar:
            self.rect(grid.margin, top + 0.12, 0.9, 0.07, theme.accent, name="Accent")
            top += 0.25
        if subtitle:
            self.text(
                grid.margin,
                top + 0.05,
                grid.content_width,
                0.45,
                subtitle,
                size=18,
                minimum=13,
                color=theme.muted,
                name="Subtitle",
                what="subtitle",
            )
            top += 0.55
        self.footer_bar()
        return max(grid.content_top, top + 0.25)

    def footer_bar(self, color: str | None = None, *, start: float | None = None) -> None:
        theme, grid = self.theme, self.grid
        color = color or theme.muted
        x = grid.margin if start is None else start
        if self.footer:
            self.text(
                x,
                grid.height - 0.5,
                grid.width - grid.margin - x - 1.2,
                0.3,
                self.footer,
                size=10,
                color=color,
                name="Footer",
            )
        if self.show_number:
            self.text(
                grid.width - grid.margin - 1.0,
                grid.height - 0.5,
                1.0,
                0.3,
                str(self.number),
                size=10,
                color=color,
                align="right",
                name="Slide Number",
            )


# -- kinds ---------------------------------------------------------------------------------


def draw(painter: Painter, slide: dict[str, Any]) -> None:
    DRAWERS[slide["kind"]](painter, slide)
    notes = slide.get("notes")
    if notes:
        painter.slide.notes_slide.notes_text_frame.text = str(notes)


def _title(painter: Painter, slide: dict[str, Any], *, closing: bool = False) -> None:
    theme, grid = painter.theme, painter.grid
    band = theme.title_background or theme.background
    text_color = theme.title_text or theme.text
    painter.background(band)
    accent = theme.palette[1] if band == theme.accent else theme.accent
    # A thin accent rule and a soft block on the right give the cover some structure.
    painter.rect(
        grid.width * 0.62,
        0,
        grid.width * 0.38,
        grid.height,
        _mix(band, text_color, 0.06),
        name="Decor",
    )
    painter.rect(
        grid.margin, grid.height * 0.30, 1.2, 0.09, readable_on(band, accent), name="Accent"
    )
    title = slide.get("title") or ("Obrigado" if closing else "")
    painter.text(
        grid.margin,
        grid.height * 0.30 + 0.3,
        grid.width * 0.72,
        2.2,
        title,
        size=48 if not closing else 54,
        minimum=28,
        font=theme.title_font,
        bold=theme.title_bold,
        color=text_color,
        anchor="top",
        name="Title",
        what="title",
    )
    below = grid.height * 0.30 + 2.6
    if slide.get("subtitle"):
        painter.text(
            grid.margin,
            below,
            grid.width * 0.72,
            0.9,
            slide["subtitle"],
            size=22,
            minimum=14,
            color=_mix(text_color, band, 0.25),
            name="Subtitle",
            what="subtitle",
        )
        below += 1.0
    details = [text for text, _ in slide.get("body") or []]
    if details:
        painter.text(
            grid.margin,
            below,
            grid.width * 0.72,
            grid.height - below - 0.6,
            "\n".join(details),
            size=16,
            minimum=11,
            color=_mix(text_color, band, 0.35),
            name="Details",
            what="details",
        )


def _section(painter: Painter, slide: dict[str, Any]) -> None:
    theme, grid = painter.theme, painter.grid
    band = theme.title_background or theme.accent
    text_color = readable_on(band, theme.title_text or theme.accent_text)
    painter.background(band)
    label = str(slide.get("number") or "")
    if label:
        painter.text(
            grid.margin,
            grid.height * 0.22,
            3,
            1.2,
            label,
            size=60,
            font=theme.title_font,
            bold=True,
            color=_mix(text_color, band, 0.45),
            name="Decor",
        )
    painter.rect(
        grid.margin,
        grid.height * 0.45,
        0.9,
        0.08,
        readable_on(band, theme.palette[1]),
        name="Accent",
    )
    painter.text(
        grid.margin,
        grid.height * 0.45 + 0.25,
        grid.content_width * 0.85,
        1.6,
        slide.get("title") or "",
        size=44,
        minimum=26,
        font=theme.title_font,
        bold=theme.title_bold,
        color=text_color,
        name="Title",
        what="title",
    )
    if slide.get("subtitle"):
        painter.text(
            grid.margin,
            grid.height * 0.45 + 1.9,
            grid.content_width * 0.85,
            0.8,
            slide["subtitle"],
            size=20,
            minimum=14,
            color=_mix(text_color, band, 0.25),
            name="Subtitle",
        )
    painter.footer_bar(_mix(text_color, band, 0.4))


def _bullets(painter: Painter, slide: dict[str, Any]) -> None:
    grid, theme = painter.grid, painter.theme
    top = painter.chrome(slide.get("title", ""), slide.get("subtitle", ""))
    body = slide.get("body") or []
    start = (
        theme.body_size + 8
        if len(body) <= 4
        else theme.body_size + 4
        if len(body) <= 7
        else theme.body_size
    )
    painter.text(
        grid.margin,
        top + 0.1,
        grid.content_width * (0.92 if len(body) > 3 else 0.85),
        grid.content_bottom - top - 0.1,
        body,
        size=start,
        minimum=12,
        bullets=True,
        gap=0.6,
        name="Body",
        what="body text",
    )
    if len(body) > 8:
        painter.warnings.append(f"slide {painter.number}: {len(body)} bullets - consider splitting")


def _column_layout(
    painter: Painter, column: dict[str, Any], w: float, h: float, card: bool
) -> tuple[float, float]:
    """(font size, box height) for a column's text."""

    theme = painter.theme
    body = column.get("body") or []
    heading_h = 0.75 if column.get("heading") else 0.0
    pad = 0.35 if card else 0.0
    measured = [(plain(t), level) for t, level in body]
    size, _ = fit.fit_size(
        measured,
        (w - 2 * pad) * 72,
        (h - heading_h - 2 * pad) * 72,
        theme.body_font,
        start=theme.body_size + 4,
        minimum=12,
        indent=0.32 * 72,
        gap=0.5,
    )
    if not card:
        return size, h
    needed = (
        fit.block_height(
            measured, (w - 2 * pad) * 72, theme.body_font, size, indent=0.32 * 72, gap=0.5
        )
        / 72
    )
    return size, min(h, max(2.4, needed + heading_h + 2 * pad + 0.3))


def _column(
    painter: Painter,
    column: dict[str, Any],
    x: float,
    top: float,
    w: float,
    h: float,
    *,
    card: str | None = None,
    size: float | None = None,
) -> None:
    theme = painter.theme
    if column.get("image"):
        painter.picture(column["image"], x, top, w, h, alt=column.get("caption", ""))
        return
    body = column.get("body") or []
    heading_h = 0.75 if column.get("heading") else 0.0
    pad = 0.35 if card else 0.0
    if size is None:
        size, h = _column_layout(painter, column, w, h, bool(card))
    if card:
        painter.rect(x, top, w, h, theme.surface, rounded=True)
        painter.rect(x, top, w, 0.1, card)
    inner_x, inner_w, inner_top, inner_h = x + pad, w - 2 * pad, top + pad, h - 2 * pad
    if column.get("heading"):
        painter.text(
            inner_x,
            inner_top,
            inner_w,
            0.6,
            column["heading"],
            size=24,
            minimum=13,
            font=theme.title_font,
            bold=True,
            color=text_safe(card or theme.accent, theme.surface if card else theme.background),
            what="column heading",
        )
        inner_top += heading_h
        inner_h -= heading_h
    painter.text(
        inner_x,
        inner_top,
        inner_w,
        inner_h,
        body,
        size=size,
        minimum=12,
        bullets=True,
        gap=0.5,
        what="column text",
    )


def _two_column(painter: Painter, slide: dict[str, Any], *, comparison: bool = False) -> None:
    grid, theme = painter.grid, painter.theme
    top = painter.chrome(slide.get("title", ""), slide.get("subtitle", ""))
    gap = 0.5
    w = (grid.content_width - gap) / 2
    h = grid.content_bottom - top
    left, right = slide.get("left") or {}, slide.get("right") or {}
    if comparison:
        # Both cards share the taller card's height and the smaller font size.
        (size_l, h_l), (size_r, h_r) = (
            _column_layout(painter, c, w, h, True) for c in (left, right)
        )
        size, h = min(size_l, size_r), max(h_l, h_r)
        _column(painter, left, grid.margin, top, w, h, card=theme.palette[0], size=size)
        _column(painter, right, grid.margin + w + gap, top, w, h, card=theme.palette[1], size=size)
        return
    _column(painter, left, grid.margin, top, w, h)
    _column(painter, right, grid.margin + w + gap, top, w, h)


def _image(painter: Painter, slide: dict[str, Any]) -> None:
    grid, theme = painter.grid, painter.theme
    if not slide.get("title"):
        painter.background(theme.background)
        painter.picture(
            slide["image"], 0, 0, grid.width, grid.height, cover=True, alt=slide.get("caption", "")
        )
        if slide.get("caption"):
            painter.rect(0, grid.height - 0.9, grid.width, 0.9, "000000")
            painter.text(
                grid.margin,
                grid.height - 0.75,
                grid.content_width,
                0.6,
                slide["caption"],
                size=18,
                minimum=12,
                color="FFFFFF",
                anchor="middle",
            )
        return
    top = painter.chrome(slide["title"], slide.get("subtitle", ""))
    caption = slide.get("caption", "")
    h = grid.content_bottom - top - (0.5 if caption else 0)
    painter.picture(slide["image"], grid.margin, top, grid.content_width, h, alt=caption)
    if caption:
        painter.text(
            grid.margin,
            top + h + 0.1,
            grid.content_width,
            0.4,
            caption,
            size=14,
            minimum=10,
            color=theme.muted,
            italic=True,
            align="center",
        )


def _image_text(painter: Painter, slide: dict[str, Any]) -> None:
    grid, theme = painter.grid, painter.theme
    painter.background(theme.background)
    image_w = grid.width * 0.46
    on_right = str(slide.get("image_side", "left")).lower() == "right"
    image_x = grid.width - image_w if on_right else 0
    painter.picture(
        slide["image"], image_x, 0, image_w, grid.height, cover=True, alt=slide.get("caption", "")
    )
    text_x = grid.margin if on_right else image_w + 0.6
    text_w = grid.width - image_w - 0.6 - grid.margin
    painter.text(
        text_x,
        0.7,
        text_w,
        1.4,
        slide.get("title", ""),
        size=theme.title_size,
        minimum=20,
        font=theme.title_font,
        bold=theme.title_bold,
        anchor="bottom",
        name="Title",
        what="title",
    )
    top = 2.25
    if theme.accent_bar:
        painter.rect(text_x, 2.2, 0.9, 0.07, theme.accent, name="Accent")
        top = 2.55
    painter.text(
        text_x,
        top,
        text_w,
        grid.height - top - 0.8,
        slide.get("body") or [],
        size=theme.body_size,
        minimum=12,
        bullets=True,
        what="body text",
    )
    painter.footer_bar(start=grid.margin if on_right else image_w + 0.6)


def _table(painter: Painter, slide: dict[str, Any]) -> None:
    grid, theme = painter.grid, painter.theme
    top = painter.chrome(slide.get("title", ""), slide.get("subtitle", ""))
    rows = [list(map(str, row)) for row in slide["table"]]
    columns = max(len(row) for row in rows)
    rows = [row + [""] * (columns - len(row)) for row in rows]
    size = 20 if len(rows) <= 5 else 16 if len(rows) <= 8 else 12 if len(rows) <= 12 else 10
    if len(rows) > 16:
        painter.warnings.append(f"slide {painter.number}: {len(rows)} table rows - split the table")
    row_h = size * 2.1 / 72
    height = min(grid.content_bottom - top, row_h * len(rows))
    weights = [max(4, max(len(plain(row[c])) for row in rows)) for c in range(columns)]
    total = sum(min(w, 40) for w in weights)
    shape = painter.slide.shapes.add_table(
        len(rows),
        columns,
        Inches(grid.margin),
        Inches(top),
        Inches(grid.content_width),
        Inches(height),
    )
    shape.name = "Table"
    table = shape.table
    tbl_pr = table._tbl.tblPr
    style = tbl_pr.find(qn("a:tableStyleId"))
    if style is not None:
        tbl_pr.remove(style)
    for c in range(columns):
        table.columns[c].width = Emu(int(Inches(grid.content_width) * min(weights[c], 40) / total))
    for r, row in enumerate(rows):
        table.rows[r].height = Inches(row_h)
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = rgb(theme.accent)
            else:
                cell.fill.fore_color.rgb = rgb(theme.surface if r % 2 == 0 else theme.background)
            frame = cell.text_frame
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            numeric = (
                plain(value)
                .replace(".", "")
                .replace(",", "")
                .replace("%", "")
                .replace("R$", "")
                .strip()
                .lstrip("-")
                .isdigit()
            )
            paragraph.alignment = PP_ALIGN.RIGHT if numeric and r else PP_ALIGN.LEFT
            for segment, seg_bold, seg_italic, code in runs(value):
                run = paragraph.add_run()
                run.text = segment
                run.font.size = Pt(size)
                run.font.name = theme.code_font if code else theme.body_font
                run.font.bold = r == 0 or seg_bold
                run.font.italic = seg_italic
                run.font.color.rgb = rgb(theme.accent_text if r == 0 else theme.text)


def _chart(painter: Painter, slide: dict[str, Any]) -> None:
    grid, theme = painter.grid, painter.theme
    top = painter.chrome(slide.get("title", ""), slide.get("subtitle", ""))
    spec = slide["chart"]
    kind = str(spec.get("type", "column")).lower()
    series = spec["series"]
    if isinstance(series, dict):
        series = [{"name": name, "values": values} for name, values in series.items()]
    if kind == "scatter":
        data = XyChartData()
        for item in series:
            points = data.add_series(str(item.get("name", "")))
            for x, y in item.get("points", []):
                points.add_data_point(float(x), float(y))
    else:
        data = CategoryChartData()
        data.categories = [str(c) for c in spec["categories"]]
        for item in series:
            values = item.get("values", [])
            if len(values) != len(spec["categories"]):
                raise ToolArgumentError(
                    f"slide {painter.number}: series {item.get('name')!r} has {len(values)} values "
                    f"for {len(spec['categories'])} categories."
                )
            data.add_series(
                str(item.get("name", "")),
                [float(v) for v in values],
                number_format=spec.get("number_format", "General"),
            )
    has_text = bool(slide.get("body"))
    width = grid.content_width * (0.64 if has_text else 1)
    frame = painter.slide.shapes.add_chart(
        CHARTS[kind],
        Inches(grid.margin),
        Inches(top),
        Inches(width),
        Inches(grid.content_bottom - top),
        data,
    )
    frame.name = "Chart"
    chart = frame.chart
    chart.font.size = Pt(13)
    chart.font.name = theme.body_font
    chart.font.color.rgb = rgb(theme.muted)
    multi = len(series) > 1 or kind in {"pie", "doughnut"}
    chart.has_legend = multi
    if multi:
        chart.legend.position = (
            XL_LEGEND_POSITION.BOTTOM
            if kind not in {"pie", "doughnut"}
            else XL_LEGEND_POSITION.RIGHT
        )
        chart.legend.include_in_layout = False
        chart.legend.font.color.rgb = rgb(theme.text)
    plot = chart.plots[0]
    if kind in {"pie", "doughnut"}:
        for index, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = rgb(theme.palette[index % len(theme.palette)])
            point.format.line.color.rgb = rgb(theme.background)
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.number_format = "0%"
        labels.number_format_is_linked = False
        labels.show_percentage = True
        labels.show_value = False
        labels.font.color.rgb = rgb("FFFFFF")
        labels.font.bold = True
        if kind == "pie":
            labels.position = XL_LABEL_POSITION.CENTER
    else:
        for index, item in enumerate(plot.series):
            color = rgb(theme.palette[index % len(theme.palette)])
            if kind in {"line", "scatter"}:
                item.format.line.color.rgb = color
                item.format.line.width = Pt(3)
                item.smooth = False
                item.marker.format.fill.solid()
                item.marker.format.fill.fore_color.rgb = color
            else:
                item.format.fill.solid()
                item.format.fill.fore_color.rgb = color
        if kind in {"column", "bar", "stacked_column", "stacked_bar"}:
            plot.gap_width = 60
            if len(series) == 1 and len(spec.get("categories", [])) <= 8:
                plot.has_data_labels = True
                plot.data_labels.font.size = Pt(12)
                plot.data_labels.font.color.rgb = rgb(theme.text)
                if kind in {"column", "bar"}:
                    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        value_axis = chart.value_axis
        value_axis.has_major_gridlines = True
        value_axis.major_gridlines.format.line.color.rgb = rgb(theme.border)
        value_axis.format.line.fill.background()
        value_axis.tick_labels.font.color.rgb = rgb(theme.muted)
        category_axis = chart.category_axis
        category_axis.format.line.color.rgb = rgb(theme.border)
        category_axis.tick_labels.font.color.rgb = rgb(theme.text)
        if spec.get("y_title"):
            value_axis.has_title = True
            value_axis.axis_title.text_frame.text = str(spec["y_title"])
    if has_text:
        x = grid.margin + width + 0.4
        painter.text(
            x,
            top + 0.2,
            grid.content_width - width - 0.4,
            grid.content_bottom - top - 0.2,
            slide["body"],
            size=18,
            minimum=12,
            bullets=True,
            what="chart notes",
        )


def _quote(painter: Painter, slide: dict[str, Any]) -> None:
    grid, theme = painter.grid, painter.theme
    painter.background(theme.surface if not theme.dark else theme.background)
    painter.text(
        grid.margin,
        0.4,
        2,
        2.8,
        "“",
        size=160,
        font="Georgia",
        color=theme.accent,
        bold=True,
        name="Decor",
    )
    painter.text(
        grid.margin + 0.6,
        1.9,
        grid.content_width - 1.2,
        3.4,
        slide["quote"],
        size=34,
        minimum=18,
        font=theme.title_font,
        italic=True,
        color=theme.text,
        anchor="middle",
        what="quote",
    )
    if slide.get("author"):
        painter.rect(grid.margin + 0.6, 5.55, 0.6, 0.06, theme.accent)
        painter.text(
            grid.margin + 0.6,
            5.75,
            grid.content_width - 1.2,
            0.5,
            slide["author"],
            size=18,
            minimum=12,
            color=theme.muted,
            bold=True,
        )
    painter.footer_bar()


def _stats(painter: Painter, slide: dict[str, Any]) -> None:
    grid, theme = painter.grid, painter.theme
    top = painter.chrome(slide.get("title", ""), slide.get("subtitle", ""))
    stats = slide["stats"]
    count = len(stats)
    per_row = count if count <= 4 else (3 if count <= 6 else 4)
    rows = (count + per_row - 1) // per_row
    gap = 0.35
    w = (grid.content_width - gap * (per_row - 1)) / per_row
    h = min(2.8, (grid.content_bottom - top - gap * (rows - 1)) / rows)
    start_y = top + max(0, (grid.content_bottom - top - rows * h - gap * (rows - 1)) / 2)
    for index, (value, label) in enumerate(stats):
        r, c = divmod(index, per_row)
        x = grid.margin + c * (w + gap)
        y = start_y + r * (h + gap)
        color = theme.palette[index % len(theme.palette)]
        painter.rect(x, y, w, h, theme.surface, rounded=True)
        painter.rect(x, y + 0.3, 0.08, h - 0.6, color)
        # One line: a big number that wraps stops reading as a number.
        # Measured at 100pt: Pillow rounds glyph advances at tiny sizes.
        widest = (fit.text_width(value, theme.title_font, 100.0, bold=True) / 100.0) or 1.0
        value_size = max(20.0, min(54 if per_row <= 3 else 44, (w - 0.7) * 72 / widest))
        painter.text(
            x + 0.35,
            y + 0.25,
            w - 0.6,
            h * 0.5,
            value,
            size=value_size,
            font=theme.title_font,
            bold=True,
            color=text_safe(color, theme.surface),
            anchor="bottom",
            what="stat value",
        )
        painter.text(
            x + 0.35,
            y + 0.3 + h * 0.5,
            w - 0.6,
            h * 0.45 - 0.35,
            label,
            size=18,
            minimum=11,
            color=theme.muted,
            what="stat label",
        )


def _timeline(painter: Painter, slide: dict[str, Any]) -> None:
    grid, theme = painter.grid, painter.theme
    top = painter.chrome(slide.get("title", ""), slide.get("subtitle", ""))
    steps = slide["steps"]
    count = len(steps)
    w = grid.content_width / count
    block = 2.9
    line_y = top + max(0.75, (grid.content_bottom - top - block) / 2 + 0.35)
    painter.rect(grid.margin + w / 2, line_y - 0.025, w * (count - 1), 0.05, theme.border)
    for index, (title, text) in enumerate(steps):
        center = grid.margin + w * index + w / 2
        color = theme.palette[index % len(theme.palette)]
        dot = painter.oval(center - 0.35, line_y - 0.35, 0.7, color)
        dot.name = f"Step {index + 1}"
        frame = dot.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = str(index + 1)
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.name = theme.title_font
        run.font.color.rgb = rgb(readable_on(color, "FFFFFF"))
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        inner = w - 0.25
        painter.text(
            center - inner / 2,
            line_y + 0.6,
            inner,
            0.8,
            title,
            size=22,
            minimum=12,
            font=theme.title_font,
            bold=True,
            align="center",
            color=theme.text,
            what="step title",
        )
        if text:
            painter.text(
                center - inner / 2,
                line_y + 1.45,
                inner,
                grid.content_bottom - line_y - 1.45,
                text,
                size=17,
                minimum=10,
                align="center",
                color=theme.muted,
                what="step text",
            )


def _cards(painter: Painter, slide: dict[str, Any]) -> None:
    grid, theme = painter.grid, painter.theme
    top = painter.chrome(slide.get("title", ""), slide.get("subtitle", ""))
    cards = slide["cards"]
    count = len(cards)
    per_row = count if count <= 3 else (2 if count == 4 else 3 if count <= 6 else 4)
    rows = (count + per_row - 1) // per_row
    gap = 0.35
    w = (grid.content_width - gap * (per_row - 1)) / per_row
    available = (grid.content_bottom - top - gap * (rows - 1)) / rows
    inner_w = (w - 0.7) * 72
    title_size = 24 if per_row <= 3 else 20
    text_size = 19 if per_row <= 3 else 16
    # Long single words (German, Portuguese) must fit a card's width unbroken.
    title_words = [word for t, _ in cards for word in plain(t).split()]
    text_words = [word for _, x in cards for word in plain(x).split()]
    while title_size > 13 and any(
        fit.text_width(word, theme.title_font, title_size, bold=True) > inner_w
        for word in title_words
    ):
        title_size -= 1
    while text_size > 10 and any(
        fit.text_width(word, theme.body_font, text_size) > inner_w for word in text_words
    ):
        text_size -= 1
    title_h = (
        max(
            fit.block_height([(plain(t), 0)], inner_w, theme.title_font, title_size, bold=True)
            for t, _ in cards
        )
        / 72
    )
    text_h = (
        max(
            fit.block_height([(plain(x), 0)], inner_w, theme.body_font, text_size) if x else 0
            for _, x in cards
        )
        / 72
    )
    h = min(available, max(2.2, 0.75 + title_h + 0.2 + text_h + 0.4))
    start_y = top + max(0.0, (grid.content_bottom - top - rows * h - gap * (rows - 1)) / 2 - 0.2)
    for index, (title, text) in enumerate(cards):
        r, c = divmod(index, per_row)
        x = grid.margin + c * (w + gap)
        y = start_y + r * (h + gap)
        color = theme.palette[index % len(theme.palette)]
        painter.rect(x, y, w, h, theme.surface, rounded=True, line=theme.border)
        painter.rect(x + 0.35, y + 0.35, 0.55, 0.08, color)
        painter.text(
            x + 0.35,
            y + 0.6,
            w - 0.7,
            title_h + 0.1,
            title,
            size=title_size,
            minimum=13,
            font=theme.title_font,
            bold=True,
            color=theme.text,
            what="card title",
        )
        if text:
            text_top = y + 0.6 + title_h + 0.25
            painter.text(
                x + 0.35,
                text_top,
                w - 0.7,
                y + h - text_top - 0.25,
                text,
                size=text_size,
                minimum=10,
                color=theme.muted,
                what="card text",
            )


def _agenda(painter: Painter, slide: dict[str, Any]) -> None:
    grid, theme = painter.grid, painter.theme
    top = painter.chrome(slide.get("title") or "Agenda", slide.get("subtitle", ""))
    items = [text for text, level in slide.get("body") or [] if level == 0]
    per_column = 5 if len(items) > 5 else len(items)
    columns = 1 if len(items) <= 5 else 2
    w = (grid.content_width - 0.5 * (columns - 1)) / columns
    row_h = min(0.95, (grid.content_bottom - top) / max(1, per_column))
    for index, item in enumerate(items[:10]):
        c, r = divmod(index, per_column)
        x = grid.margin + c * (w + 0.5)
        y = top + r * row_h
        painter.text(
            x,
            y,
            0.9,
            row_h,
            f"{index + 1:02d}",
            size=28,
            font=theme.title_font,
            bold=True,
            color=text_safe(theme.palette[index % len(theme.palette)], theme.background),
            anchor="middle",
        )
        painter.text(
            x + 1.0,
            y,
            w - 1.0,
            row_h,
            item,
            size=22,
            minimum=13,
            color=theme.text,
            anchor="middle",
            what="agenda item",
        )
        if r < per_column - 1 and index < len(items) - 1:
            painter.rect(x + 1.0, y + row_h - 0.02, w - 1.0, 0.015, theme.border)


def _code(painter: Painter, slide: dict[str, Any]) -> None:
    grid, theme = painter.grid, painter.theme
    top = painter.chrome(slide.get("title", ""), slide.get("subtitle", ""))
    has_text = bool(slide.get("body"))
    width = grid.content_width * (0.64 if has_text else 1)
    h = grid.content_bottom - top
    painter.rect(grid.margin, top, width, h, theme.code_background, rounded=True, name="Code Panel")
    for index, color in enumerate(("FF5F57", "FEBC2E", "28C840")):
        painter.oval(grid.margin + 0.25 + index * 0.25, top + 0.2, 0.14, color)
    lines = slide["code"].split("\n")
    longest = max(lines, key=len) if lines else ""
    size = 20.0
    while size > 9 and (
        fit.text_width(longest, theme.code_font, size) > (width - 0.6) * 72
        or len(lines) * size * 1.25 > (h - 0.7) * 72
    ):
        size -= 0.5
    if size <= 9:
        painter.warnings.append(f"slide {painter.number}: code is too long for one slide")
    box = painter.slide.shapes.add_textbox(
        Inches(grid.margin + 0.3), Inches(top + 0.5), Inches(width - 0.6), Inches(h - 0.7)
    )
    box.name = "Code"
    frame = box.text_frame
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.line_spacing = 1.1
        run = paragraph.add_run()
        run.text = line.replace("\t", "    ") or " "
        run.font.name = theme.code_font
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(_code_color(line, theme))
    if has_text:
        x = grid.margin + width + 0.4
        painter.text(
            x,
            top,
            grid.content_width - width - 0.4,
            h,
            slide["body"],
            size=18,
            minimum=12,
            bullets=True,
            what="code notes",
        )


def _code_color(line: str, theme: Theme) -> str:
    stripped = line.strip()
    if stripped.startswith(("#", "//", "--", "/*", "*")):
        return "7F8EA3"
    return theme.code_text


def _closing(painter: Painter, slide: dict[str, Any]) -> None:
    _title(painter, slide, closing=True)


def _mix(a: str, b: str, amount: float) -> str:
    """``a`` moved ``amount`` of the way towards ``b``."""

    a_rgb = [int(a[i : i + 2], 16) for i in (0, 2, 4)]
    b_rgb = [int(b[i : i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * amount):02X}" for x, y in zip(a_rgb, b_rgb, strict=True))


DRAWERS: dict[str, Callable[[Painter, dict[str, Any]], None]] = {
    "title": _title,
    "section": _section,
    "bullets": _bullets,
    "two_column": _two_column,
    "comparison": lambda p, s: _two_column(p, s, comparison=True),
    "image": _image,
    "image_text": _image_text,
    "table": _table,
    "chart": _chart,
    "quote": _quote,
    "stats": _stats,
    "timeline": _timeline,
    "cards": _cards,
    "agenda": _agenda,
    "code": _code,
    "closing": _closing,
}

_unused = MSO_CONNECTOR
