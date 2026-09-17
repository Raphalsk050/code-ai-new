"""A deck's content and the problems a reviewer would point at."""

from __future__ import annotations

from collections import Counter
from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Pt

from code_ai.tools.slides import fit
from code_ai.tools.slides.themes import contrast

EMU = 914400
MIN_BODY_PT = 12
_CHROME_NAMES = {"footer", "slide number", "decor", "accent", "details"}
# Titles lower than this sit on cover or divider slides, which are meant to differ.
_CONTENT_TITLE_TOP = 1.5 * EMU


def inches(value) -> float:
    return round((value or 0) / EMU, 2)


def shape_kind(shape) -> str:
    if shape.is_placeholder:
        return "placeholder"
    if shape.has_chart if hasattr(shape, "has_chart") else False:
        return "chart"
    if shape.has_table if hasattr(shape, "has_table") else False:
        return "table"
    kind = shape.shape_type
    if kind == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if kind == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if shape.has_text_frame and shape.text_frame.text.strip():
        return "text"
    return "shape"


def is_title(shape) -> bool:
    if shape.is_placeholder:
        try:
            return shape.placeholder_format.type in (
                PP_PLACEHOLDER.TITLE,
                PP_PLACEHOLDER.CENTER_TITLE,
            )
        except ValueError:
            return False
    return shape.name.strip().lower() == "title"


def slide_title(slide) -> str:
    for shape in slide.shapes:
        if is_title(shape) and shape.has_text_frame:
            return shape.text_frame.text.strip()
    return ""


def background_hex(slide) -> str:
    for source in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        try:
            fill = source.background.fill
            if fill.type == 1:  # solid
                return str(fill.fore_color.rgb)
        except (AttributeError, TypeError, ValueError):
            continue
    return "FFFFFF"


def _fill_hex(shape) -> str | None:
    try:
        if shape.fill.type == 1:
            return str(shape.fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError):
        return None
    return None


def _run_color(run) -> str | None:
    try:
        if run.font.color is not None and run.font.color.type is not None:
            return str(run.font.color.rgb)
    except (AttributeError, TypeError, ValueError):
        return None
    return None


def describe_shape(shape, text_chars: int) -> dict[str, Any]:
    entry: dict[str, Any] = {
        "id": shape.shape_id,
        "name": shape.name,
        "kind": shape_kind(shape),
        "box_in": [
            inches(shape.left),
            inches(shape.top),
            inches(shape.width),
            inches(shape.height),
        ],
    }
    if shape.has_text_frame and shape.text_frame.text.strip():
        text = shape.text_frame.text.strip()
        entry["text"] = text[:text_chars] + ("..." if len(text) > text_chars else "")
        sizes = sorted(
            {r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size}
        )
        fonts = sorted(
            {r.font.name for p in shape.text_frame.paragraphs for r in p.runs if r.font.name}
        )
        if sizes:
            entry["font_sizes"] = sizes
        if fonts:
            entry["fonts"] = fonts
    if entry["kind"] == "table":
        table = shape.table
        entry["table"] = f"{len(table.rows)}x{len(table.columns)}"
        entry["header"] = [cell.text for cell in table.rows[0].cells][:8]
    if entry["kind"] == "chart":
        chart = shape.chart
        entry["chart"] = str(chart.chart_type).split(".")[-1].split(" ")[0]
        try:
            entry["categories"] = [str(c) for c in chart.plots[0].categories][:12]
            entry["series"] = {s.name: list(s.values)[:12] for s in chart.plots[0].series}
        except (IndexError, AttributeError, TypeError):
            pass
    if entry["kind"] == "picture":
        entry["alt_text"] = shape._element.nvPicPr.cNvPr.get("descr") or ""
    if shape.is_placeholder:
        try:
            entry["placeholder"] = str(shape.placeholder_format.type).split(".")[-1].split(" ")[0]
        except ValueError:
            pass
    return entry


def inspect_deck(deck, *, slides: list[int], text_chars: int) -> dict[str, Any]:
    fonts: Counter[str] = Counter()
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name and run.text.strip():
                            fonts[run.font.name] += len(run.text)
    report: dict[str, Any] = {
        "slides": len(deck.slides),
        "size_in": [inches(deck.slide_width), inches(deck.slide_height)],
        "fonts_in_use": [f"{name} ({chars} chars)" for name, chars in fonts.most_common(10)],
        "layouts_in_use": dict(Counter(slide.slide_layout.name for slide in deck.slides)),
    }
    listed = []
    for index in slides:
        slide = deck.slides[index]
        entry: dict[str, Any] = {
            "slide": index + 1,
            "layout": slide.slide_layout.name,
            "title": slide_title(slide),
            "shapes": [describe_shape(shape, text_chars) for shape in slide.shapes],
        }
        if slide._element.get("show") == "0":
            entry["hidden"] = True
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            entry["notes"] = slide.notes_slide.notes_text_frame.text.strip()[:400]
        listed.append(entry)
    report["slide_details"] = listed
    report["issues"] = find_issues(deck)
    return report


def estimated_overflow(shape) -> float:
    """How many points of text do not fit the box (0 when it fits or cannot be judged)."""

    if not shape.has_text_frame or not shape.text_frame.text.strip() or not shape.width:
        return 0.0
    frame = shape.text_frame
    if frame.word_wrap is False:
        return 0.0
    width = (shape.width - frame.margin_left - frame.margin_right) / EMU * 72
    height = (shape.height - frame.margin_top - frame.margin_bottom) / EMU * 72
    total = 0.0
    for number, paragraph in enumerate(frame.paragraphs):
        runs = [run for run in paragraph.runs if run.text]
        text = "".join(run.text for run in runs)
        size = next((run.font.size.pt for run in runs if run.font.size), 18.0)
        family = next((run.font.name for run in runs if run.font.name), "Calibri")
        bold = any(run.font.bold for run in runs)
        indent = (paragraph.level + 1) * 0.3 * 72 if paragraph.level else 0
        lines = len(fit.wrap(text, max(20.0, width - indent), family, size, bold)) if text else 1
        total += lines * size * fit.LINE_SPACING
        if number and paragraph.space_before is not None:
            total += paragraph.space_before.pt
    return max(0.0, total - height)


def find_issues(deck) -> list[dict[str, Any]]:
    issues: dict[str, dict[str, Any]] = {}
    slide_w, slide_h = deck.slide_width, deck.slide_height

    def flag(rule: str, message: str, where: str) -> None:
        entry = issues.setdefault(rule, {"rule": rule, "message": message, "count": 0, "at": []})
        entry["count"] += 1
        if len(entry["at"]) < 10:
            entry["at"].append(where)

    title_boxes: Counter[tuple] = Counter()
    title_fonts: Counter[tuple] = Counter()
    for number, slide in enumerate(deck.slides, start=1):
        background = background_hex(slide)
        if not slide_title(slide):
            flag(
                "missing-title",
                "Slide without a title (screen readers and navigation need one)",
                f"slide {number}",
            )
        words = 0
        bullets = 0
        text_boxes = []
        for shape in slide.shapes:
            where = f"slide {number} shape {shape.shape_id} ({shape.name})"
            if (
                shape.left is not None
                and (
                    shape.left < -Pt(2)
                    or shape.top < -Pt(2)
                    or shape.left + shape.width > slide_w + Pt(2)
                    or shape.top + shape.height > slide_h + Pt(2)
                )
                and shape_kind(shape) != "picture"
            ):
                flag("off-slide", "Shape extends past the slide edge", where)
            if shape.is_placeholder and shape.has_text_frame and not shape.text_frame.text.strip():
                flag(
                    "empty-placeholder",
                    "Empty placeholder ('Click to add text' shows in edit mode)",
                    where,
                )
            if (
                shape_kind(shape) == "picture"
                and not (shape._element.nvPicPr.cNvPr.get("descr") or "").strip()
            ):
                flag("image-alt-text", "Picture without alternative text", where)
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            name = shape.name.strip().lower()
            if name in _CHROME_NAMES:
                continue
            text_boxes.append(shape)
            if is_title(shape) and (shape.top or 0) < _CONTENT_TITLE_TOP:
                title_boxes[(inches(shape.left), inches(shape.top))] += 1
                size = next(
                    (
                        r.font.size.pt
                        for p in shape.text_frame.paragraphs
                        for r in p.runs
                        if r.font.size
                    ),
                    None,
                )
                family = next(
                    (
                        r.font.name
                        for p in shape.text_frame.paragraphs
                        for r in p.runs
                        if r.font.name
                    ),
                    None,
                )
                title_fonts[(family, size)] += 1
            elif not is_title(shape):
                words += len(shape.text_frame.text.split())
                points = sum(1 for p in shape.text_frame.paragraphs if p.text.strip())
                bullets += points if points > 1 else 0
            overflow = estimated_overflow(shape)
            if overflow > 6:
                flag(
                    "text-overflow", f"Text likely overflows its box (by ~{overflow:.0f}pt)", where
                )
            if shape.name:
                small = [
                    r.font.size.pt
                    for p in shape.text_frame.paragraphs
                    for r in p.runs
                    if r.font.size and r.font.size.pt < MIN_BODY_PT and r.text.strip()
                ]
                if small:
                    flag(
                        "small-text",
                        f"Text below {MIN_BODY_PT}pt (smallest {min(small):g}pt)",
                        where,
                    )
            fill = _fill_hex(shape) or _underlying_fill(slide, shape) or background
            for paragraph in shape.text_frame.paragraphs:
                colors = {c for c in (_run_color(r) for r in paragraph.runs if r.text.strip()) if c}
                low = [c for c in colors if contrast(c, fill) < 3.0]
                if low:
                    flag("low-contrast", f"Text colour #{low[0]} on #{fill} is hard to read", where)
                    break
        if words > 90:
            flag("wordy-slide", f"Over 90 words on one slide ({words})", f"slide {number}")
        if bullets > 9:
            flag("too-many-points", f"{bullets} text points on one slide", f"slide {number}")
        for i, first in enumerate(text_boxes):
            for second in text_boxes[i + 1 :]:
                if _overlap(first, second) > 0.2:
                    flag(
                        "overlapping-text",
                        "Text boxes overlap",
                        f"slide {number} shapes {first.shape_id}/{second.shape_id}",
                    )
    if len(title_boxes) > 1 and sum(title_boxes.values()) > 2:
        common = title_boxes.most_common(1)[0][0]
        outliers = sum(count for box, count in title_boxes.items() if box != common)
        # One deliberate exception (an image-beside-text slide) is design, not drift.
        if outliers >= 2 or outliers / sum(title_boxes.values()) >= 0.25:
            flag(
                "title-position",
                f"Titles sit in {len(title_boxes)} different places",
                f"{outliers} slide(s) off the usual {common}",
            )
    if len(title_fonts) > 2:
        flag(
            "title-style",
            "Titles use inconsistent fonts or sizes",
            ", ".join(f"{f} {s}" for (f, s) in list(title_fonts)[:4]),
        )
    return sorted(issues.values(), key=lambda item: -item["count"])


def _underlying_fill(slide, shape) -> str | None:
    """Fill of the topmost filled shape drawn before ``shape`` that contains its centre."""

    cx = (shape.left or 0) + (shape.width or 0) / 2
    cy = (shape.top or 0) + (shape.height or 0) / 2
    found = None
    for other in slide.shapes:
        if other.shape_id == shape.shape_id:
            break
        if other.left is None or other.width is None:
            continue
        inside = (
            other.left <= cx <= other.left + other.width
            and other.top <= cy <= other.top + other.height
        )
        colour = _fill_hex(other) if inside else None
        if colour:
            found = colour
    return found


def _overlap(a, b) -> float:
    ax1, ay1 = a.left + a.width, a.top + a.height
    bx1, by1 = b.left + b.width, b.top + b.height
    width = min(ax1, bx1) - max(a.left, b.left)
    height = min(ay1, by1) - max(a.top, b.top)
    if width <= 0 or height <= 0:
        return 0.0
    smaller = min(a.width * a.height, b.width * b.height) or 1
    return width * height / smaller
