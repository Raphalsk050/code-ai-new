"""slides_edit's operations. Slides are 1-based; shapes are ids or names from slides_inspect."""

from __future__ import annotations

import copy
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from code_ai.core.errors import ToolArgumentError
from code_ai.tools.office.common import parse_page_spec
from code_ai.tools.slides import deck as decks
from code_ai.tools.slides import spec as specs
from code_ai.tools.slides.themes import Theme, rgb


@dataclass
class DeckState:
    deck: Any
    theme: Theme
    resolve_image: Callable[[str], Path]
    footer: str = ""
    warnings: list[str] = field(default_factory=list)


def fail(index: int, message: str) -> ToolArgumentError:
    return ToolArgumentError(f"operation {index}: {message}")


def _slide_index(state: DeckState, value: Any, index: int) -> int:
    try:
        return parse_page_spec(value, len(state.deck.slides))[0]
    except (ToolArgumentError, IndexError) as exc:
        raise fail(index, f"slide {value!r}: {exc}") from None


def _slides(state: DeckState, value: Any, index: int) -> list[int]:
    try:
        return parse_page_spec(value, len(state.deck.slides))
    except ToolArgumentError as exc:
        raise fail(index, str(exc)) from None


def _shape(state: DeckState, spec: dict[str, Any], index: int):
    slide = state.deck.slides[_slide_index(state, spec.get("slide"), index)]
    wanted = spec.get("shape")
    if wanted is None:
        raise fail(index, "'shape' (id or name from slides_inspect) is required.")
    for shape in slide.shapes:
        if (
            str(shape.shape_id) == str(wanted)
            or shape.name.strip().lower() == str(wanted).strip().lower()
        ):
            return slide, shape
    names = ", ".join(f"{s.shape_id}:{s.name}" for s in slide.shapes)
    raise fail(index, f"no shape {wanted!r} on that slide. Shapes: {names}.")


# -- slides ----------------------------------------------------------------------------------


def op_add(state: DeckState, spec: dict[str, Any], index: int) -> str:
    content = spec.get("slide") if spec.get("slide") is not None else spec.get("markdown")
    if content is None:
        content = {k: v for k, v in spec.items() if k not in {"op", "at"}}
    parsed = specs.parse(content if isinstance(content, str | list) else [content])
    at = spec.get("at")
    position = None if at in (None, "end") else _slide_index(state, at, index)
    added = 0
    for offset, slide_spec in enumerate(parsed):
        number = (position if position is not None else len(state.deck.slides)) + offset + 1
        state.warnings += decks.add_slide(
            state.deck,
            slide_spec,
            state.theme,
            resolve_image=state.resolve_image,
            number=number,
            footer=state.footer,
            position=None if position is None else position + offset,
        )
        added += 1
    return f"added {added} slide(s)"


def op_duplicate(state: DeckState, spec: dict[str, Any], index: int) -> str:
    source_index = _slide_index(state, spec.get("slide"), index)
    duplicate_slide(state.deck, source_index)
    target = spec.get("at")
    if target is not None:
        decks.move_slide(state.deck, len(state.deck.slides) - 1, _slide_index(state, target, index))
    else:
        decks.move_slide(state.deck, len(state.deck.slides) - 1, source_index + 1)
    return f"duplicated slide {source_index + 1}"


def duplicate_slide(deck, source_index: int):
    """Copy a slide with its pictures, charts and notes; charts get their own parts."""

    source = deck.slides[source_index]
    copy_slide = deck.slides.add_slide(source.slide_layout)
    for shape in list(copy_slide.shapes):
        shape._element.getparent().remove(shape._element)
    rid_map: dict[str, str] = {}
    for rel in source.part.rels.values():
        if rel.reltype == RT.NOTES_SLIDE or rel.reltype == RT.SLIDE_LAYOUT:
            continue
        if rel.is_external:
            rid_map[rel.rId] = copy_slide.part.relate_to(
                rel.target_ref, rel.reltype, is_external=True
            )
        elif rel.reltype == RT.CHART:
            rid_map[rel.rId] = _copy_chart(copy_slide, rel.target_part)
        else:
            rid_map[rel.rId] = copy_slide.part.relate_to(rel.target_part, rel.reltype)
    tree = copy_slide.shapes._spTree
    for element in source.shapes._spTree.iterchildren():
        if element.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        clone = copy.deepcopy(element)
        for node in clone.iter():
            for attribute in (qn("r:embed"), qn("r:link"), qn("r:id")):
                value = node.get(attribute)
                if value in rid_map:
                    node.set(attribute, rid_map[value])
        tree.append(clone)
    background = source._element.cSld.find(qn("p:bg"))
    if background is not None:
        copy_slide._element.cSld.insert(0, copy.deepcopy(background))
    if source.has_notes_slide:
        copy_slide.notes_slide.notes_text_frame.text = source.notes_slide.notes_text_frame.text
    return copy_slide


def _copy_chart(slide, chart_part) -> str:
    """A new chart part with the original's data and formatting."""

    chart = chart_part.chart
    data = CategoryChartData()
    try:
        plot = chart.plots[0]
        data.categories = list(plot.categories)
        for series in plot.series:
            data.add_series(series.name, list(series.values))
        chart_type = chart.chart_type
    except Exception:  # noqa: BLE001 - exotic charts: fall back to sharing the part
        return slide.part.relate_to(chart_part, RT.CHART)
    frame = slide.shapes.add_chart(chart_type, 0, 0, Inches(1), Inches(1), data)
    # Resolved through the relationship, not frame.chart: that caches a Chart bound to the
    # XML about to be replaced, and later edits would go to the discarded copy.
    rel_id = frame._element.xpath(".//c:chart/@r:id")[0]
    new_part = slide.part.related_part(rel_id)
    workbook_rel = new_part._element.find(qn("c:externalData"))
    workbook_id = workbook_rel.get(qn("r:id")) if workbook_rel is not None else None
    styled = copy.deepcopy(chart_part._element)
    external = styled.find(qn("c:externalData"))
    if external is not None and workbook_id:
        external.set(qn("r:id"), workbook_id)
    new_part._element = styled
    new_part.__dict__.pop("chart", None)
    frame._element.getparent().remove(frame._element)
    return rel_id


def op_delete(state: DeckState, spec: dict[str, Any], index: int) -> str:
    doomed = sorted(set(_slides(state, spec.get("slides", spec.get("slide")), index)), reverse=True)
    if len(doomed) >= len(state.deck.slides):
        raise fail(index, "deleting every slide leaves an empty deck.")
    slide_ids = state.deck.slides._sldIdLst
    for position in doomed:
        element = list(slide_ids)[position]
        state.deck.part.drop_rel(element.rId)
        slide_ids.remove(element)
    return f"deleted {len(doomed)} slide(s)"


def op_move(state: DeckState, spec: dict[str, Any], index: int) -> str:
    source = _slide_index(state, spec.get("slide"), index)
    target = _slide_index(state, spec.get("to"), index)
    decks.move_slide(state.deck, source, target)
    return f"moved slide {source + 1} to position {target + 1}"


def op_hide(state: DeckState, spec: dict[str, Any], index: int) -> str:
    chosen = _slides(state, spec.get("slides", spec.get("slide")), index)
    hidden = spec.get("op") == "hide"
    for position in chosen:
        element = state.deck.slides[position]._element
        if hidden:
            element.set("show", "0")
        elif "show" in element.attrib:
            del element.attrib["show"]
    return f"{'hid' if hidden else 'unhid'} {len(chosen)} slide(s)"


# -- text ------------------------------------------------------------------------------------


def op_replace_text(state: DeckState, spec: dict[str, Any], index: int) -> str:
    find = spec.get("find")
    if not isinstance(find, str) or not find:
        raise fail(index, "'find' is required.")
    replacement = str(spec.get("replace", ""))
    flags = 0 if spec.get("match_case", True) else re.IGNORECASE
    pattern = re.compile(find if spec.get("regex") else re.escape(find), flags)
    if not spec.get("regex"):
        replacement = replacement.replace("\\", "\\\\")
    chosen = _slides(state, spec.get("slides"), index)
    count = 0
    for position in chosen:
        slide = state.deck.slides[position]
        frames = [shape.text_frame for shape in _all_shapes(slide.shapes) if shape.has_text_frame]
        for shape in _all_shapes(slide.shapes):
            if getattr(shape, "has_table", False) and shape.has_table:
                frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
        if slide.has_notes_slide and spec.get("include_notes"):
            frames.append(slide.notes_slide.notes_text_frame)
        for frame in frames:
            for paragraph in frame.paragraphs:
                count += _replace_runs(paragraph, pattern, replacement)
    return f"replaced {count} occurrence(s) of {find!r}"


def _all_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:  # group
            yield from _all_shapes(shape.shapes)


def _replace_runs(paragraph, pattern: re.Pattern, replacement: str) -> int:
    done = 0
    cursor = 0
    while True:
        runs = paragraph.runs
        text = "".join(run.text for run in runs)
        match = pattern.search(text, cursor)
        if match is None:
            return done
        new_text = match.expand(replacement)
        position = 0
        placed = False
        for run in runs:
            run_text = run.text
            start, end = position, position + len(run_text)
            position = end
            if not placed and (match.start() < end or (match.start() == end and run is runs[-1])):
                run.text = (
                    run_text[: match.start() - start]
                    + new_text
                    + run_text[min(match.end(), end) - start :]
                )
                placed = True
                if match.end() <= end:
                    break
            elif placed:
                if start >= match.end():
                    break
                run.text = run_text[min(match.end(), end) - start :]
        cursor = match.start() + len(new_text) + (1 if match.start() == match.end() else 0)
        done += 1


def op_set_text(state: DeckState, spec: dict[str, Any], index: int) -> str:
    _, shape = _shape(state, spec, index)
    if not shape.has_text_frame:
        raise fail(index, f"shape {shape.name!r} holds no text.")
    frame = shape.text_frame
    template = None
    for paragraph in frame.paragraphs:
        if paragraph.runs:
            template = paragraph.runs[0]._r.find(qn("a:rPr"))
            break
    paragraph_template = frame.paragraphs[0]._p.find(qn("a:pPr"))
    lines = specs._markdown_body(str(spec.get("text", "")).split("\n")) or [("", 0)]
    body = frame._txBody
    for paragraph in list(frame.paragraphs)[1:]:
        body.remove(paragraph._p)
    first = frame.paragraphs[0]
    for run in list(first.runs):
        first._p.remove(run._r)
    for number, (text, level) in enumerate(lines):
        paragraph = first if number == 0 else frame.add_paragraph()
        if number and paragraph_template is not None:
            paragraph._p.insert(0, copy.deepcopy(paragraph_template))
        paragraph.level = min(level, 8)
        for segment, bold, italic, _code in specs.runs(text):
            run = paragraph.add_run()
            run.text = segment
            if template is not None:
                run._r.insert(0, copy.deepcopy(template))
            if bold:
                run.font.bold = True
            if italic:
                run.font.italic = True
    if spec.get("size"):
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(float(spec["size"]))
    return f"set the text of {shape.name!r}"


def op_add_text(state: DeckState, spec: dict[str, Any], index: int) -> str:
    slide = state.deck.slides[_slide_index(state, spec.get("slide"), index)]
    from code_ai.tools.slides.layouts import Grid, Painter

    painter = Painter(slide, state.theme, Grid(0, 0), state.resolve_image, 0)
    box = painter.text(
        _inch(spec, "x", 1.0),
        _inch(spec, "y", 1.0),
        _inch(spec, "width", 6.0),
        _inch(spec, "height", 1.0),
        specs._markdown_body(str(spec.get("text", "")).split("\n")) or [("", 0)],
        size=float(spec.get("size") or 18),
        color=spec.get("color", "").lstrip("#") or None,
        bold=bool(spec.get("bold")),
        align=str(spec.get("align") or "left"),
        bullets=bool(spec.get("bullets")),
        name=str(spec.get("name") or "Text"),
    )
    return f"added text box {box.shape_id}"


def op_add_image(state: DeckState, spec: dict[str, Any], index: int) -> str:
    slide = state.deck.slides[_slide_index(state, spec.get("slide"), index)]
    if not spec.get("image"):
        raise fail(index, "'image' is required.")
    from code_ai.tools.slides.layouts import Grid, Painter

    painter = Painter(slide, state.theme, Grid(0, 0), state.resolve_image, 0)
    picture = painter.picture(
        str(spec["image"]),
        _inch(spec, "x", 1.0),
        _inch(spec, "y", 1.0),
        _inch(spec, "width", 5.0),
        _inch(spec, "height", 3.0),
        cover=bool(spec.get("cover")),
        alt=str(spec.get("alt") or ""),
    )
    state.warnings += painter.warnings
    return f"added picture {picture.shape_id}"


def _inch(spec: dict[str, Any], key: str, default: float) -> float:
    value = spec.get(key)
    if value is None:
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        raise ToolArgumentError(f"'{key}' is a number of inches.") from None


def op_move_shape(state: DeckState, spec: dict[str, Any], index: int) -> str:
    _, shape = _shape(state, spec, index)
    for key, attribute in (("x", "left"), ("y", "top"), ("width", "width"), ("height", "height")):
        if spec.get(key) is not None:
            setattr(shape, attribute, Inches(_inch(spec, key, 0)))
    return f"moved/resized {shape.name!r}"


def op_delete_shape(state: DeckState, spec: dict[str, Any], index: int) -> str:
    _, shape = _shape(state, spec, index)
    shape._element.getparent().remove(shape._element)
    return f"deleted shape {shape.name!r}"


def op_alt_text(state: DeckState, spec: dict[str, Any], index: int) -> str:
    _, shape = _shape(state, spec, index)
    shape._element.xpath("./*[1]/p:cNvPr")[0].set("descr", str(spec.get("text", "")))
    return f"set alt text of {shape.name!r}"


def op_notes(state: DeckState, spec: dict[str, Any], index: int) -> str:
    slide = state.deck.slides[_slide_index(state, spec.get("slide"), index)]
    slide.notes_slide.notes_text_frame.text = str(spec.get("text", ""))
    return "set speaker notes"


def op_background(state: DeckState, spec: dict[str, Any], index: int) -> str:
    chosen = _slides(state, spec.get("slides", spec.get("slide")), index)
    for position in chosen:
        slide = state.deck.slides[position]
        if spec.get("image"):
            from code_ai.tools.slides.layouts import Grid, Painter

            grid = decks.grid_for(state.deck)
            painter = Painter(
                slide, state.theme, Grid(grid.width, grid.height), state.resolve_image, position + 1
            )
            picture = painter.picture(
                str(spec["image"]), 0, 0, grid.width, grid.height, cover=True, alt="background"
            )
            tree = slide.shapes._spTree
            tree.remove(picture._element)
            tree.insert(2, picture._element)
        elif spec.get("color"):
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = rgb(str(spec["color"]))
        else:
            raise fail(index, "background needs 'color' or 'image'.")
    return f"set the background of {len(chosen)} slide(s)"


def op_table_cell(state: DeckState, spec: dict[str, Any], index: int) -> str:
    slide = state.deck.slides[_slide_index(state, spec.get("slide"), index)]
    table_shape = _pick(slide, spec, "has_table", index, "table")
    table = table_shape.table
    try:
        cell = table.cell(int(spec["row"]), int(spec["column"]))
    except (KeyError, IndexError, ValueError, TypeError):
        raise fail(
            index, f"row/column out of range for a {len(table.rows)}x{len(table.columns)} table."
        ) from None
    paragraph = cell.text_frame.paragraphs[0]
    template = paragraph.runs[0]._r.find(qn("a:rPr")) if paragraph.runs else None
    for extra in list(cell.text_frame.paragraphs)[1:]:
        cell.text_frame._txBody.remove(extra._p)
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run()
    run.text = str(spec.get("text", ""))
    if template is not None:
        run._r.insert(0, copy.deepcopy(template))
    return f"set cell ({spec['row']}, {spec['column']})"


def op_update_chart(state: DeckState, spec: dict[str, Any], index: int) -> str:
    slide = state.deck.slides[_slide_index(state, spec.get("slide"), index)]
    chart = _pick(slide, spec, "has_chart", index, "chart").chart
    categories = spec.get("categories")
    series = spec.get("series")
    if isinstance(series, dict):
        series = [{"name": name, "values": values} for name, values in series.items()]
    if not categories or not isinstance(series, list) or not series:
        raise fail(index, 'update_chart needs categories and series, e.g. {"Revenue": [1, 2]}.')
    data = CategoryChartData()
    data.categories = [str(c) for c in categories]
    for item in series:
        if len(item.get("values", [])) != len(categories):
            raise fail(index, f"series {item.get('name')!r} needs {len(categories)} values.")
        data.add_series(str(item.get("name", "")), [float(v) for v in item["values"]])
    chart.replace_data(data)
    return "updated chart data"


def _pick(slide, spec: dict[str, Any], attribute: str, index: int, what: str):
    candidates = [s for s in slide.shapes if getattr(s, attribute, False)]
    if spec.get("shape") is not None:
        candidates = [
            s
            for s in candidates
            if str(s.shape_id) == str(spec["shape"]) or s.name.lower() == str(spec["shape"]).lower()
        ]
    if not candidates:
        raise fail(index, f"no {what} found on that slide.")
    return candidates[0]


HANDLERS: dict[str, Callable[[DeckState, dict[str, Any], int], str]] = {
    "add": op_add,
    "duplicate": op_duplicate,
    "delete": op_delete,
    "move": op_move,
    "hide": op_hide,
    "unhide": op_hide,
    "replace_text": op_replace_text,
    "set_text": op_set_text,
    "add_text": op_add_text,
    "add_image": op_add_image,
    "move_shape": op_move_shape,
    "delete_shape": op_delete_shape,
    "alt_text": op_alt_text,
    "notes": op_notes,
    "background": op_background,
    "table_cell": op_table_cell,
    "update_chart": op_update_chart,
}

ALIASES = {
    "add_slide": "add",
    "insert": "add",
    "duplicate_slide": "duplicate",
    "copy": "duplicate",
    "delete_slide": "delete",
    "remove": "delete",
    "move_slide": "move",
    "reorder": "move",
    "replace": "replace_text",
    "edit_text": "set_text",
    "text": "set_text",
    "textbox": "add_text",
    "image": "add_image",
    "resize": "move_shape",
    "remove_shape": "delete_shape",
    "speaker_notes": "notes",
    "set_notes": "notes",
    "set_background": "background",
    "chart_data": "update_chart",
    "cell": "table_cell",
}


def normalize(operations: Any) -> list[dict[str, Any]]:
    if isinstance(operations, dict):
        operations = [operations]
    if not isinstance(operations, list) or not operations:
        raise ToolArgumentError("'operations' must be a non-empty JSON array of objects.")
    result = []
    for index, raw in enumerate(operations, start=1):
        if not isinstance(raw, dict):
            raise ToolArgumentError(f"operation {index} must be an object with an 'op' key.")
        name = str(raw.get("op") or "").strip().lower()
        canonical = ALIASES.get(name, name)
        if canonical not in HANDLERS:
            raise ToolArgumentError(
                f"operation {index}: unknown op {raw.get('op')!r}. Known: {sorted(HANDLERS)}."
            )
        result.append({**raw, "op": canonical})
    return result


def apply(state: DeckState, operations: list[dict[str, Any]]) -> list[str]:
    return [
        f"{i}. {HANDLERS[spec['op']](state, spec, i)}" for i, spec in enumerate(operations, start=1)
    ]
