"""WordprocessingML that python-docx has no API for: fields, list numbering, links, borders."""

from __future__ import annotations

import copy

from docx.document import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph

BULLET_CHARS = ("•", "◦", "▪", "•", "◦", "▪", "•", "◦", "▪")


def _el(tag: str, **attrs: str):
    element = OxmlElement(tag)
    for key, value in attrs.items():
        element.set(qn(key), str(value))
    return element


def add_field(paragraph: Paragraph, instruction: str, placeholder: str = "1"):
    """A complex field (PAGE, NUMPAGES, TOC...), shown with ``placeholder`` until updated."""

    run = paragraph.add_run()
    run._r.append(_el("w:fldChar", **{"w:fldCharType": "begin", "w:dirty": "true"}))
    instr = paragraph.add_run()
    text = _el("w:instrText", **{"xml:space": "preserve"})
    text.text = f" {instruction} "
    instr._r.append(text)
    separate = paragraph.add_run()
    separate._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    shown = paragraph.add_run(placeholder)
    end = paragraph.add_run()
    end._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
    return shown


def update_fields_on_open(document: Document) -> None:
    """Ask Word to refresh TOC and page fields when the file is opened."""

    settings = document.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = _el("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def add_hyperlink(paragraph: Paragraph, url: str, text: str, *, color: str = "0B5CAD"):
    part = paragraph.part
    if url.startswith("#"):
        link = _el("w:hyperlink", **{"w:anchor": url[1:]})
    else:
        rel_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
        link = _el("w:hyperlink", **{"r:id": rel_id})
    run = _el("w:r")
    props = _el("w:rPr")
    props.append(_el("w:color", **{"w:val": color}))
    props.append(_el("w:u", **{"w:val": "single"}))
    run.append(props)
    text_el = _el("w:t", **{"xml:space": "preserve"})
    text_el.text = text
    run.append(text_el)
    link.append(run)
    paragraph._p.append(link)
    return link


def shade(element_pr, fill: str) -> None:
    """Background fill on a paragraph (pPr) or cell (tcPr) properties element."""

    for old in element_pr.findall(qn("w:shd")):
        element_pr.remove(old)
    element_pr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))


def paragraph_border(paragraph: Paragraph, side: str = "bottom", color: str = "BFBFBF", size=6):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = _el("w:pBdr")
        p_pr.append(borders)
    borders.append(
        _el(f"w:{side}", **{"w:val": "single", "w:sz": size, "w:space": 4, "w:color": color})
    )


def table_borders(table, color: str = "BFBFBF", size: int = 4) -> None:
    tbl_pr = table._tbl.tblPr
    for old in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(old)
    borders = _el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(
            _el(f"w:{side}", **{"w:val": "single", "w:sz": size, "w:space": 0, "w:color": color})
        )
    tbl_pr.append(borders)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(_el("w:tblHeader", **{"w:val": "true"}))


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(_el("w:cantSplit", **{"w:val": "true"}))


# -- numbering -------------------------------------------------------------------------


class Numbering:
    """Bullet and decimal list definitions owned by Code-AI, one w:num per ordered list.

    The template's "List Number" style shares one counter across the whole
    document, so a second list would continue at 4. Each ordered list here gets
    its own w:num, restarting at 1.
    """

    def __init__(self, document: Document) -> None:
        self._numbering = document.part.numbering_part.element
        self._bullet_abstract = self._abstract("bullet")
        self._decimal_abstract = self._abstract("decimal")
        self._bullet_num: int | None = None

    def bullet_list(self) -> int:
        if self._bullet_num is None:
            self._bullet_num = self._num(self._bullet_abstract)
        return self._bullet_num

    def ordered_list(self, start: int = 1) -> int:
        return self._num(self._decimal_abstract, restart=start)

    def heading_numbering(self, levels: int = 3) -> int:
        """Outline numbering 1, 1.1, 1.1.1 linked to Heading 1..n."""

        abstract_id = self._next_id("w:abstractNum", "w:abstractNumId")
        abstract = _el("w:abstractNum", **{"w:abstractNumId": abstract_id})
        abstract.append(_el("w:multiLevelType", **{"w:val": "multilevel"}))
        for level in range(9):
            lvl = _el("w:lvl", **{"w:ilvl": level})
            lvl.append(_el("w:start", **{"w:val": 1}))
            lvl.append(_el("w:numFmt", **{"w:val": "decimal" if level < levels else "none"}))
            if level < levels:
                lvl.append(_el("w:pStyle", **{"w:val": f"Heading{level + 1}"}))
            text = ".".join(f"%{i + 1}" for i in range(level + 1)) if level < levels else ""
            lvl.append(_el("w:lvlText", **{"w:val": text}))
            lvl.append(_el("w:lvlJc", **{"w:val": "left"}))
            p_pr = _el("w:pPr")
            p_pr.append(_el("w:ind", **{"w:left": 0, "w:firstLine": 0}))
            lvl.append(p_pr)
            lvl.append(_el("w:suff", **{"w:val": "space"}))
            abstract.append(lvl)
        self._insert_abstract(abstract)
        return self._num(abstract_id)

    def _abstract(self, kind: str) -> int:
        abstract_id = self._next_id("w:abstractNum", "w:abstractNumId")
        abstract = _el("w:abstractNum", **{"w:abstractNumId": abstract_id})
        abstract.append(_el("w:multiLevelType", **{"w:val": "hybridMultilevel"}))
        for level in range(9):
            lvl = _el("w:lvl", **{"w:ilvl": level})
            lvl.append(_el("w:start", **{"w:val": 1}))
            if kind == "bullet":
                lvl.append(_el("w:numFmt", **{"w:val": "bullet"}))
                lvl.append(_el("w:lvlText", **{"w:val": BULLET_CHARS[level]}))
            else:
                fmt = ("decimal", "lowerLetter", "lowerRoman")[level % 3]
                lvl.append(_el("w:numFmt", **{"w:val": fmt}))
                lvl.append(_el("w:lvlText", **{"w:val": f"%{level + 1}."}))
            lvl.append(_el("w:lvlJc", **{"w:val": "left"}))
            p_pr = _el("w:pPr")
            p_pr.append(_el("w:ind", **{"w:left": 360 * (level + 1) + 360, "w:hanging": 360}))
            lvl.append(p_pr)
            abstract.append(lvl)
        self._insert_abstract(abstract)
        return abstract_id

    def _insert_abstract(self, abstract) -> None:
        # Schema order: every abstractNum before the first num.
        first_num = self._numbering.find(qn("w:num"))
        if first_num is None:
            self._numbering.append(abstract)
        else:
            first_num.addprevious(abstract)

    def _num(self, abstract_id: int, restart: int | None = None) -> int:
        num_id = self._next_id("w:num", "w:numId")
        num = _el("w:num", **{"w:numId": num_id})
        num.append(_el("w:abstractNumId", **{"w:val": abstract_id}))
        if restart is not None:
            override = _el("w:lvlOverride", **{"w:ilvl": 0})
            override.append(_el("w:startOverride", **{"w:val": restart}))
            num.append(override)
        self._numbering.append(num)
        return num_id

    def _next_id(self, tag: str, attribute: str) -> int:
        used = [int(el.get(qn(attribute))) for el in self._numbering.findall(qn(tag))]
        return max(used, default=0) + 1


def set_numbering(paragraph: Paragraph, num_id: int, level: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for old in p_pr.findall(qn("w:numPr")):
        p_pr.remove(old)
    num_pr = _el("w:numPr")
    num_pr.append(_el("w:ilvl", **{"w:val": max(0, min(8, level))}))
    num_pr.append(_el("w:numId", **{"w:val": num_id}))
    # numPr sits right after pStyle, or first.
    style = p_pr.find(qn("w:pStyle"))
    if style is not None:
        style.addnext(num_pr)
    else:
        p_pr.insert(0, num_pr)


def link_style_numbering(style, num_id: int, level: int) -> None:
    p_pr = style.element.get_or_add_pPr()
    for old in p_pr.findall(qn("w:numPr")):
        p_pr.remove(old)
    num_pr = _el("w:numPr")
    num_pr.append(_el("w:ilvl", **{"w:val": level}))
    num_pr.append(_el("w:numId", **{"w:val": num_id}))
    p_pr.insert(0, num_pr)


def list_info(paragraph: Paragraph) -> tuple[str, int] | None:
    """('bullet' | 'number', level) when the paragraph is a list item, directly or by style."""

    p_pr = paragraph._p.pPr
    num_pr = p_pr.find(qn("w:numPr")) if p_pr is not None else None
    style_name = paragraph.style.name if paragraph.style is not None else ""
    if num_pr is None:
        if style_name.startswith("List Bullet"):
            return "bullet", _style_level(style_name)
        if style_name.startswith("List Number"):
            return "number", _style_level(style_name)
        return None
    num_id_el = num_pr.find(qn("w:numId"))
    ilvl_el = num_pr.find(qn("w:ilvl"))
    if num_id_el is None or num_id_el.get(qn("w:val")) == "0":
        return None
    level = int(ilvl_el.get(qn("w:val"))) if ilvl_el is not None else 0
    fmt = _num_format(paragraph, int(num_id_el.get(qn("w:val"))), level)
    return ("bullet" if fmt == "bullet" else "number"), level


def _style_level(name: str) -> int:
    tail = name.split()[-1]
    return int(tail) - 1 if tail.isdigit() else 0


def _num_format(paragraph: Paragraph, num_id: int, level: int) -> str:
    try:
        numbering = paragraph.part.numbering_part.element
    except (AttributeError, KeyError, NotImplementedError):
        return "decimal"
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) != str(num_id):
            continue
        abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        for abstract in numbering.findall(qn("w:abstractNum")):
            if abstract.get(qn("w:abstractNumId")) != abstract_id:
                continue
            for lvl in abstract.findall(qn("w:lvl")):
                if lvl.get(qn("w:ilvl")) == str(level):
                    fmt = lvl.find(qn("w:numFmt"))
                    return fmt.get(qn("w:val")) if fmt is not None else "decimal"
    return "decimal"


def hex_color(value: str) -> RGBColor:
    text = value.strip().lstrip("#")
    if len(text) == 3:
        text = "".join(ch * 2 for ch in text)
    return RGBColor.from_string(text.upper())


def clone(element):
    return copy.deepcopy(element)


def set_language(document: Document, tag: str) -> None:
    """Default proofing language, e.g. pt-BR, in docDefaults and Normal."""

    styles = document.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    targets = []
    if defaults is not None:
        r_pr_default = defaults.find(qn("w:rPrDefault"))
        if r_pr_default is not None and r_pr_default.find(qn("w:rPr")) is not None:
            targets.append(r_pr_default.find(qn("w:rPr")))
    targets.append(document.styles["Normal"].element.get_or_add_rPr())
    for r_pr in targets:
        for old in r_pr.findall(qn("w:lang")):
            r_pr.remove(old)
        r_pr.append(_el("w:lang", **{"w:val": tag, "w:eastAsia": tag, "w:bidi": tag}))
