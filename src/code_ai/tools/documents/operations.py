"""document_edit's operations. Indexes are the p/t numbers document_inspect prints."""

from __future__ import annotations

import copy
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx.document import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

from code_ai.core.errors import ToolArgumentError
from code_ai.tools.documents import builder, ooxml, presets
from code_ai.tools.documents.inspect import heading_level

ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

HIGHLIGHTS = {
    "yellow": WD_COLOR_INDEX.YELLOW,
    "green": WD_COLOR_INDEX.BRIGHT_GREEN,
    "cyan": WD_COLOR_INDEX.TURQUOISE,
    "pink": WD_COLOR_INDEX.PINK,
    "gray": WD_COLOR_INDEX.GRAY_25,
    "red": WD_COLOR_INDEX.RED,
    "none": None,
}


@dataclass
class DocState:
    document: Document
    resolve_image: Callable[[str], Path]
    language: str
    notes: list[str] = field(default_factory=list)

    @property
    def paragraphs(self) -> list[Paragraph]:
        return self.document.paragraphs


def fail(index: int, message: str) -> ToolArgumentError:
    return ToolArgumentError(f"operation {index}: {message}")


def indexes(spec: Any, total: int, index: int, what: str = "paragraph") -> list[int]:
    """Zero-based indexes, as document_inspect numbers them: 3, '3-7,9', [1, 2]."""

    if isinstance(spec, int):
        parts = [str(spec)]
    elif isinstance(spec, list):
        parts = [str(item) for item in spec]
    elif isinstance(spec, str):
        parts = spec.replace(" ", "").lstrip("pt").split(",")
    else:
        raise fail(index, f"'{what}s' must be an index, a range like '3-7' or a list.")
    result: list[int] = []
    for part in parts:
        part = part.lstrip("pt")
        if not part:
            continue
        try:
            if "-" in part:
                start, end = (int(x.lstrip("pt")) for x in part.split("-", 1))
                result.extend(range(start, end + 1))
            else:
                result.append(int(part))
        except ValueError:
            raise fail(index, f"unreadable {what} index {part!r}.") from None
    for value in result:
        if not 0 <= value < total:
            raise fail(index, f"{what} {value} does not exist (there are {total}).")
    return result


def find_heading(state: DocState, text: str, index: int) -> int:
    wanted = text.strip().casefold()
    candidates = [
        position
        for position, paragraph in enumerate(state.paragraphs)
        if heading_level(paragraph) is not None and wanted in paragraph.text.strip().casefold()
    ]
    if not candidates:
        raise fail(index, f"no heading contains {text!r}.")
    exact = [p for p in candidates if state.paragraphs[p].text.strip().casefold() == wanted]
    return (exact or candidates)[0]


def section_end(state: DocState, heading_index: int) -> int:
    """Index of the last paragraph belonging to the heading's section."""

    paragraphs = state.paragraphs
    level = heading_level(paragraphs[heading_index]) or 1
    end = len(paragraphs) - 1
    for position in range(heading_index + 1, len(paragraphs)):
        other = heading_level(paragraphs[position])
        if other is not None and other <= level:
            end = position - 1
            break
    # Trailing spacers and page breaks belong before the next section, not inside this one.
    while end > heading_index and _is_blank(paragraphs[end]):
        end -= 1
    return end


def _is_blank(paragraph: Paragraph) -> bool:
    return not paragraph.text.strip() and paragraph._p.find(f".//{qn('w:drawing')}") is None


# -- text across runs ------------------------------------------------------------------


def _pattern(spec: dict[str, Any], index: int) -> re.Pattern:
    find = spec.get("find")
    if not isinstance(find, str) or not find:
        raise fail(index, "'find' is required.")
    flags = 0 if spec.get("match_case", True) else re.IGNORECASE
    source = find if spec.get("regex") else re.escape(find)
    if spec.get("whole_word"):
        source = rf"\b{source}\b"
    try:
        return re.compile(source, flags)
    except re.error as exc:
        raise fail(index, f"invalid regex: {exc}") from None


def _all_paragraphs(document: Document, scope: str):
    yield from document.paragraphs
    if scope == "body":
        return
    for table in document.tables:
        yield from _table_paragraphs(table)
    seen: list = []
    for section in document.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
        ):
            if part.is_linked_to_previous or any(part._element is other for other in seen):
                continue
            seen.append(part._element)
            yield from part.paragraphs
            for table in part.tables:
                yield from _table_paragraphs(table)


def _table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for inner in cell.tables:
                yield from _table_paragraphs(inner)


def replace_in_paragraph(
    paragraph: Paragraph, pattern: re.Pattern, replacement: str, limit: int
) -> int:
    """Replace matches that may span runs; the replacement takes the first run's formatting."""

    done = 0
    cursor = 0
    while limit <= 0 or done < limit:
        runs = paragraph.runs
        text = "".join(run.text for run in runs)
        match = pattern.search(text, cursor)
        if match is None:
            break
        new_text = match.expand(replacement)
        _splice(runs, match.start(), match.end(), new_text)
        cursor = match.start() + len(new_text) + (1 if match.start() == match.end() else 0)
        done += 1
        if cursor > len(text) + len(new_text):
            break
    return done


def _splice(runs, start: int, end: int, new_text: str) -> None:
    position = 0
    placed = False
    for number, run in enumerate(runs):
        text = run.text
        run_start, run_end = position, position + len(text)
        position = run_end
        if not placed:
            if start < run_end or (start == run_end and number == len(runs) - 1):
                local_start = start - run_start
                local_end = min(end, run_end) - run_start
                run.text = text[:local_start] + new_text + text[local_end:]
                placed = True
                if end <= run_end:
                    return
        else:
            if run_start >= end:
                return
            run.text = text[min(end, run_end) - run_start :]
            if end <= run_end:
                return


def isolate(paragraph: Paragraph, start: int, end: int) -> list:
    """Split runs so [start, end) is covered by whole runs, and return those runs."""

    covered = []
    position = 0
    for run in list(paragraph.runs):
        text = run.text
        run_start, run_end = position, position + len(text)
        position = run_end
        if run_end <= start or run_start >= end:
            continue
        cut_start = max(start, run_start) - run_start
        cut_end = min(end, run_end) - run_start
        if cut_end < len(text):
            tail = copy.deepcopy(run._r)
            run._r.addnext(tail)
            _set_run_text(tail, text[cut_end:])
            run.text = text[:cut_end]
            text = text[:cut_end]
        if cut_start > 0:
            head = copy.deepcopy(run._r)
            run._r.addprevious(head)
            _set_run_text(head, text[:cut_start])
            run.text = text[cut_start:]
        covered.append(run)
    return covered


def _set_run_text(r_element, text: str) -> None:
    from docx.text.run import Run

    Run(r_element, None).text = text


# -- operations ------------------------------------------------------------------------


def op_replace(state: DocState, spec: dict[str, Any], index: int) -> str:
    pattern = _pattern(spec, index)
    replacement = spec.get("replace")
    if replacement is None:
        raise fail(index, "'replace' is required (use \"\" to delete the text).")
    replacement = str(replacement)
    if not spec.get("regex"):
        replacement = replacement.replace("\\", "\\\\")
    limit = int(spec.get("count") or 0)
    scope = str(spec.get("scope") or "all").lower()
    chosen = None
    if spec.get("paragraphs") is not None:
        # Keep the elements referenced so their ids stay theirs while we compare.
        kept = [
            state.paragraphs[i]._p
            for i in indexes(spec["paragraphs"], len(state.paragraphs), index)
        ]
        chosen = {id(element) for element in kept}
    total = 0
    for paragraph in _all_paragraphs(state.document, scope):
        if chosen is not None and id(paragraph._p) not in chosen:
            continue
        remaining = limit - total if limit else 0
        if limit and remaining <= 0:
            break
        total += replace_in_paragraph(paragraph, pattern, replacement, remaining)
    if not total:
        state.notes.append(f"operation {index}: {spec['find']!r} was not found.")
    return f"replaced {total} occurrence(s) of {spec['find']!r}"


def op_format_text(state: DocState, spec: dict[str, Any], index: int) -> str:
    pattern = _pattern(spec, index)
    count = 0
    for paragraph in _all_paragraphs(state.document, str(spec.get("scope") or "body")):
        text = "".join(run.text for run in paragraph.runs)
        for match in reversed(list(pattern.finditer(text))):
            if match.start() == match.end():
                continue
            for run in isolate(paragraph, match.start(), match.end()):
                _apply_run_format(run, spec, index)
            count += 1
    return f"formatted {count} occurrence(s) of {spec['find']!r}"


def _apply_run_format(run, spec: dict[str, Any], index: int) -> None:
    for key in ("bold", "italic", "underline"):
        if spec.get(key) is not None:
            setattr(run, key, bool(spec[key]))
    if spec.get("strike") is not None:
        run.font.strike = bool(spec["strike"])
    if spec.get("color"):
        run.font.color.rgb = ooxml.hex_color(str(spec["color"]))
    if spec.get("font"):
        run.font.name = str(spec["font"])
    if spec.get("size"):
        run.font.size = Pt(float(spec["size"]))
    if spec.get("highlight") is not None:
        key = str(spec["highlight"]).lower()
        if key not in HIGHLIGHTS:
            raise fail(index, f"highlight is one of {sorted(HIGHLIGHTS)}.")
        run.font.highlight_color = HIGHLIGHTS[key]


def _anchor(state: DocState, spec: dict[str, Any], index: int) -> tuple[Any, str]:
    """The body element to insert next to, and 'after' or 'before'."""

    body = state.document.element.body
    paragraphs = state.paragraphs
    if spec.get("after_heading"):
        heading = find_heading(state, str(spec["after_heading"]), index)
        return paragraphs[section_end(state, heading)]._p, "after"
    if spec.get("after") is not None:
        target = indexes(spec["after"], len(paragraphs), index)[0]
        return paragraphs[target]._p, "after"
    if spec.get("before") is not None:
        target = indexes(spec["before"], len(paragraphs), index)[0]
        return paragraphs[target]._p, "before"
    if spec.get("after_table") is not None:
        table = indexes(spec["after_table"], len(state.document.tables), index, "table")[0]
        return state.document.tables[table]._tbl, "after"
    where = str(spec.get("at") or "end").lower()
    children = [child for child in body if child.tag != qn("w:sectPr")]
    if where == "start" and children:
        return children[0], "before"
    return (children[-1], "after") if children else (None, "after")


def _render_at(
    state: DocState, spec: dict[str, Any], index: int, render: Callable[[builder.DocxWriter], None]
) -> int:
    anchor, side = _anchor(state, spec, index)
    body = state.document.element.body
    # Position, not id(): lxml hands out fresh proxy objects for the same element.
    before = len(_content(body))
    presets.ensure_styles(state.document)
    writer = builder.DocxWriter(
        state.document,
        presets.PRESETS["default"],
        resolve_image=state.resolve_image,
        language=state.language,
    )
    render(writer)
    state.notes.extend(writer.notes)
    created = _content(body)[before:]
    if anchor is None:
        return len(created)
    if side == "after":
        for element in created:
            anchor.addnext(element)
            anchor = element
    else:
        for element in created:
            anchor.addprevious(element)
    return len(created)


def _content(body) -> list:
    return [child for child in body if child.tag != qn("w:sectPr")]


def op_insert(state: DocState, spec: dict[str, Any], index: int) -> str:
    markdown = spec.get("markdown") if spec.get("markdown") is not None else spec.get("text")
    if not isinstance(markdown, str) or not markdown.strip():
        raise fail(index, "'markdown' (the content to insert) is required.")
    count = _render_at(
        state,
        spec,
        index,
        lambda writer: writer.write_markdown(markdown, title_from_single_h1=False),
    )
    return f"inserted {count} block(s)"


def op_delete(state: DocState, spec: dict[str, Any], index: int) -> str:
    if spec.get("section"):
        heading = find_heading(state, str(spec["section"]), index)
        doomed = list(range(heading, section_end(state, heading) + 1))
    elif spec.get("paragraphs") is not None:
        doomed = indexes(spec["paragraphs"], len(state.paragraphs), index)
    elif spec.get("table") is not None:
        table = indexes(spec["table"], len(state.document.tables), index, "table")[0]
        element = state.document.tables[table]._tbl
        element.getparent().remove(element)
        return f"deleted table {table}"
    else:
        raise fail(index, "delete needs 'paragraphs', 'section' or 'table'.")
    elements = [state.paragraphs[i]._p for i in sorted(set(doomed))]
    for element in elements:
        element.getparent().remove(element)
    return f"deleted {len(elements)} paragraph(s)"


def op_set_style(state: DocState, spec: dict[str, Any], index: int) -> str:
    style = str(spec.get("style") or "")
    presets.ensure_styles(state.document)
    try:
        state.document.styles[style]
    except KeyError:
        names = sorted(s.name for s in state.document.styles if s.type == 1)
        raise fail(
            index, f"no paragraph style {style!r}. Styles: {', '.join(names[:40])}."
        ) from None
    targets = indexes(spec.get("paragraphs"), len(state.paragraphs), index)
    for target in targets:
        state.paragraphs[target].style = state.document.styles[style]
    return f"set style {style!r} on {len(targets)} paragraph(s)"


def op_paragraph_format(state: DocState, spec: dict[str, Any], index: int) -> str:
    targets = indexes(spec.get("paragraphs"), len(state.paragraphs), index)
    for target in targets:
        fmt = state.paragraphs[target].paragraph_format
        if spec.get("alignment"):
            key = str(spec["alignment"]).lower()
            if key not in ALIGNMENTS:
                raise fail(index, f"alignment is one of {sorted(ALIGNMENTS)}.")
            fmt.alignment = ALIGNMENTS[key]
        for key in ("space_before", "space_after"):
            if spec.get(key) is not None:
                setattr(fmt, key, Pt(float(spec[key])))
        if spec.get("line_spacing") is not None:
            fmt.line_spacing = float(spec["line_spacing"])
        for key in ("first_line_indent", "left_indent", "right_indent"):
            if spec.get(key) is not None:
                setattr(fmt, key, Cm(float(str(spec[key]).rstrip("cm"))))
        for key in ("keep_with_next", "page_break_before", "keep_together"):
            if spec.get(key) is not None:
                setattr(fmt, key, bool(spec[key]))
    return f"formatted {len(targets)} paragraph(s)"


def _table(state: DocState, spec: dict[str, Any], index: int):
    tables = state.document.tables
    number = indexes(spec.get("table", 0), len(tables), index, "table")[0]
    return tables[number]


def op_table_cell(state: DocState, spec: dict[str, Any], index: int) -> str:
    table = _table(state, spec, index)
    try:
        cell = table.cell(int(spec["row"]), int(spec["column"]))
    except (KeyError, IndexError, ValueError, TypeError):
        raise fail(
            index, f"row/column out of range: the table is {len(table.rows)}x{len(table.columns)}."
        ) from None
    paragraph = cell.paragraphs[0]
    runs = paragraph.runs
    template = copy.deepcopy(runs[0]._r) if runs else None
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    for run in runs:
        run._r.getparent().remove(run._r)
    run = paragraph.add_run(str(spec.get("text", "")))
    if template is not None:
        r_pr = template.find(qn("w:rPr"))
        if r_pr is not None:
            run._r.insert(0, r_pr)
    if spec.get("bold") is not None:
        run.bold = bool(spec["bold"])
    return f"set cell ({spec['row']}, {spec['column']})"


def op_table_add_row(state: DocState, spec: dict[str, Any], index: int) -> str:
    table = _table(state, spec, index)
    values = spec.get("values")
    if not isinstance(values, list):
        raise fail(index, "'values' is a list of cell texts.")
    reference = table.rows[-1]._tr
    row = copy.deepcopy(reference)
    if spec.get("after") is not None:
        position = int(spec["after"])
        if not 0 <= position < len(table.rows):
            raise fail(index, "'after' is a row index of the table.")
        table.rows[position]._tr.addnext(row)
    else:
        reference.addnext(row)
    from docx.table import _Row

    new_row = _Row(row, table)
    for column, cell in enumerate(new_row.cells):
        for extra in cell.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        paragraph = cell.paragraphs[0]
        runs = paragraph.runs
        for run in runs[1:]:
            run._r.getparent().remove(run._r)
        text = str(values[column]) if column < len(values) else ""
        if runs:
            runs[0].text = text
        else:
            paragraph.add_run(text)
    return f"added a row to the table ({len(table.rows)} rows now)"


def op_table_delete_row(state: DocState, spec: dict[str, Any], index: int) -> str:
    table = _table(state, spec, index)
    row = int(spec.get("row", -1))
    if not 0 <= row < len(table.rows):
        raise fail(index, f"row {row} does not exist.")
    element = table.rows[row]._tr
    element.getparent().remove(element)
    return f"deleted row {row}"


def op_page_break(state: DocState, spec: dict[str, Any], index: int) -> str:
    count = _render_at(state, spec, index, lambda writer: writer.page_break())
    return f"inserted a page break ({count} block)"


def op_toc(state: DocState, spec: dict[str, Any], index: int) -> str:
    levels = int(spec.get("levels") or 3)
    _render_at(state, spec, index, lambda writer: writer.toc(levels))
    return "inserted a table of contents (Word fills it in when the file is opened)"


def op_header_footer(state: DocState, spec: dict[str, Any], index: int) -> str:
    builder.set_header_footer(
        state.document,
        presets.PRESETS["default"] if not spec.get("position") else _positioned(spec["position"]),
        header=spec.get("header"),
        footer=spec.get("footer"),
        page_numbers=bool(spec.get("page_numbers")),
        title=state.document.core_properties.title or "",
        skip_first_page=bool(spec.get("skip_first_page")),
    )
    return "updated header/footer"


def _positioned(position: str) -> presets.Preset:
    from dataclasses import replace

    return replace(presets.PRESETS["default"], page_number_position=str(position).lower())


def op_properties(state: DocState, spec: dict[str, Any], index: int) -> str:
    core = state.document.core_properties
    changed = []
    for key in ("title", "author", "subject", "keywords", "comments", "category"):
        if spec.get(key) is not None:
            setattr(core, key, str(spec[key]))
            changed.append(key)
    if not changed:
        raise fail(index, "give at least one of title, author, subject, keywords, comments.")
    return f"set properties: {', '.join(changed)}"


def op_page_setup(state: DocState, spec: dict[str, Any], index: int) -> str:
    preset = presets.PRESETS["default"]
    section = state.document.sections[0]
    current = tuple(
        round(value / 360000, 3)
        for value in (
            section.top_margin,
            section.right_margin,
            section.bottom_margin,
            section.left_margin,
        )
    )
    presets.apply_page_setup(
        state.document,
        preset,
        page_size=spec.get("page_size") or _size_name(section),
        orientation=spec.get("orientation")
        or ("landscape" if section.page_width > section.page_height else "portrait"),
        margins=presets.parse_margins(spec.get("margins")) or current,
    )
    return "updated page setup"


def _size_name(section) -> str:
    width = min(section.page_width, section.page_height) / 360000
    height = max(section.page_width, section.page_height) / 360000
    for name, (w, h) in presets.PAGE_SIZES_CM.items():
        if abs(w - width) < 0.2 and abs(h - height) < 0.2:
            return name
    return "a4"


def op_comment(state: DocState, spec: dict[str, Any], index: int) -> str:
    target = indexes(spec.get("paragraph"), len(state.paragraphs), index)[0]
    paragraph = state.paragraphs[target]
    runs = paragraph.runs
    if spec.get("find"):
        text = "".join(run.text for run in runs)
        position = text.find(str(spec["find"]))
        if position < 0:
            raise fail(index, f"{spec['find']!r} is not in paragraph {target}.")
        runs = isolate(paragraph, position, position + len(str(spec["find"])))
    if not runs:
        runs = [paragraph.add_run("")]
    state.document.add_comment(
        runs, text=str(spec.get("text") or ""), author=str(spec.get("author") or "Code-AI")
    )
    return f"commented on paragraph {target}"


def op_accept_changes(state: DocState, spec: dict[str, Any], index: int) -> str:
    return _resolve_revisions(state.document, accept=True)


def op_reject_changes(state: DocState, spec: dict[str, Any], index: int) -> str:
    return _resolve_revisions(state.document, accept=False)


def _resolve_revisions(document: Document, *, accept: bool) -> str:
    root = document.element
    keep, drop = (
        (("w:ins", "w:moveTo"), ("w:del", "w:moveFrom"))
        if accept
        else (
            ("w:del", "w:moveFrom"),
            ("w:ins", "w:moveTo"),
        )
    )
    changed = 0
    for tag in drop:
        for element in list(root.iter(qn(tag))):
            parent = element.getparent()
            if parent is not None and parent.tag != qn("w:rPr"):
                parent.remove(element)
                changed += 1
    for tag in keep:
        for element in list(root.iter(qn(tag))):
            parent = element.getparent()
            if parent is None or parent.tag == qn("w:rPr"):
                continue
            for child in list(element):
                element.addprevious(child)
            parent.remove(element)
            changed += 1
    for text in list(root.iter(qn("w:delText"))):
        text.tag = qn("w:t")
    for tag in ("w:rPrChange", "w:pPrChange", "w:sectPrChange", "w:tblPrChange", "w:trPrChange"):
        for element in list(root.iter(qn(tag))):
            element.getparent().remove(element)
    for marker in ("w:ins", "w:del"):
        for element in list(root.iter(qn(marker))):
            if element.getparent() is not None and element.getparent().tag == qn("w:rPr"):
                element.getparent().remove(element)
    return f"{'accepted' if accept else 'rejected'} {changed} tracked change(s)"


HANDLERS = {
    "replace": op_replace,
    "format_text": op_format_text,
    "insert": op_insert,
    "delete": op_delete,
    "set_style": op_set_style,
    "paragraph_format": op_paragraph_format,
    "table_cell": op_table_cell,
    "table_add_row": op_table_add_row,
    "table_delete_row": op_table_delete_row,
    "page_break": op_page_break,
    "toc": op_toc,
    "header_footer": op_header_footer,
    "properties": op_properties,
    "page_setup": op_page_setup,
    "comment": op_comment,
    "accept_changes": op_accept_changes,
    "reject_changes": op_reject_changes,
}

ALIASES = {
    "find_replace": "replace",
    "insert_markdown": "insert",
    "add": "insert",
    "remove": "delete",
    "style": "set_style",
    "format": "format_text",
    "set_cell": "table_cell",
    "add_row": "table_add_row",
    "delete_row": "table_delete_row",
    "metadata": "properties",
    "add_comment": "comment",
    "header": "header_footer",
    "footer": "header_footer",
}


def normalize(operations: Any) -> list[dict[str, Any]]:
    if isinstance(operations, dict):
        operations = [operations]
    if not isinstance(operations, list) or not operations:
        raise ToolArgumentError("'operations' must be a non-empty JSON array of objects.")
    result = []
    for index, raw in enumerate(operations, start=1):
        if not isinstance(raw, dict):
            raise ToolArgumentError(f"operation {index} must be an object with an 'op' key.")
        name = str(raw.get("op") or "").strip().lower()
        name = ALIASES.get(name, name)
        if name not in HANDLERS:
            raise ToolArgumentError(
                f"operation {index}: unknown op {raw.get('op')!r}. Known: {sorted(HANDLERS)}."
            )
        result.append({**raw, "op": name})
    return result


def apply(state: DocState, operations: list[dict[str, Any]]) -> list[str]:
    done = []
    for index, spec in enumerate(operations, start=1):
        done.append(f"{index}. {HANDLERS[spec['op']](state, spec, index)}")
    return done


def page_break_run(paragraph: Paragraph) -> None:
    paragraph.add_run().add_break(WD_BREAK.PAGE)
