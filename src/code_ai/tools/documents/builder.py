"""Markdown into a python-docx Document, through real styles, lists, fields and captions."""

from __future__ import annotations

import re
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx.document import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

from code_ai.tools.documents import ooxml
from code_ai.tools.documents.presets import (
    CODE_STYLE,
    INLINE_CODE_STYLE,
    LONG_QUOTE_STYLE,
    Preset,
    table_text_width_cm,
)

_PAGEBREAK = re.compile(r"^\s*(\\newpage|\\pagebreak|<!--\s*pagebreak\s*-->)\s*$", re.IGNORECASE)
_TOC = re.compile(r"^\s*(\[toc\]|<!--\s*toc\s*-->)\s*$", re.IGNORECASE)
_ATTRS = re.compile(r"^\{\s*:?\s*([^}]*)\}")
_TABLE_CAPTION = re.compile(r"^(table|tabela|quadro)\s*:\s*(.+)$", re.IGNORECASE)

LABELS = {
    "pt": {"figure": "Figura", "table": "Tabela", "toc": "Sumário", "toc_hint": "Atualize o campo"},
    "en": {"figure": "Figure", "table": "Table", "toc": "Contents", "toc_hint": "Update field"},
}


def labels_for(language: str) -> dict[str, str]:
    return LABELS["pt" if language.lower().startswith("pt") else "en"]


@dataclass
class _ListContext:
    ordered: bool
    num_id: int
    level: int


@dataclass
class DocxWriter:
    document: Document
    preset: Preset
    resolve_image: Callable[[str], Path]
    language: str = "en"
    notes: list[str] = field(default_factory=list)
    numbering: ooxml.Numbering | None = None
    headings: list[tuple[int, str]] = field(default_factory=list)
    _figures: int = 0
    _tables: int = 0
    _title_h1: bool = False
    _shift: int = 0

    def __post_init__(self) -> None:
        if self.numbering is None:
            self.numbering = ooxml.Numbering(self.document)
        self.labels = labels_for(self.language)

    # -- entry points --------------------------------------------------------------

    def write_markdown(self, text: str, *, title_from_single_h1: bool = True) -> None:
        from markdown_it import MarkdownIt

        parser = MarkdownIt("commonmark", {"html": True, "typographer": True})
        parser.enable(["table", "strikethrough"])
        lines = []
        for line in text.splitlines():
            if _PAGEBREAK.match(line):
                lines.append("\n<!-- pagebreak -->\n")
            elif _TOC.match(line):
                lines.append("\n<!-- toc -->\n")
            else:
                lines.append(line)
        tokens = parser.parse("\n".join(lines))
        h1_count = sum(1 for t in tokens if t.type == "heading_open" and t.tag == "h1")
        self._title_h1 = bool(
            title_from_single_h1
            and h1_count == 1
            and tokens
            and tokens[0].type == "heading_open"
            and tokens[0].tag == "h1"
        )
        # With the lone H1 as the title, '##' sections are the document's first level.
        self._shift = 1 if self._title_h1 else 0
        self._walk(tokens)

    # -- block walker ----------------------------------------------------------------

    def _walk(self, tokens) -> None:
        lists: list[_ListContext] = []
        quote_depth = 0
        pending_caption: str | None = None
        item_started = False
        index = 0
        while index < len(tokens):
            token = tokens[index]
            kind = token.type
            if kind == "heading_open":
                level = int(token.tag[1])
                inline = tokens[index + 1]
                if level == 1 and self._title_h1:
                    paragraph = self.document.add_paragraph(style="Title")
                    self._title_h1 = False
                else:
                    level = max(1, level - self._shift)
                    paragraph = self.document.add_paragraph(style=f"Heading {min(level, 6)}")
                    self.headings.append((level, inline.content))
                self._inline(paragraph, inline.children or [])
                index += 3
                continue
            if kind in {"bullet_list_open", "ordered_list_open"}:
                ordered = kind == "ordered_list_open"
                if ordered:
                    start = int(token.attrGet("start") or 1)
                    num_id = self.numbering.ordered_list(start)
                else:
                    num_id = self.numbering.bullet_list()
                lists.append(_ListContext(ordered, num_id, len(lists)))
            elif kind in {"bullet_list_close", "ordered_list_close"}:
                lists.pop()
            elif kind == "list_item_open":
                item_started = True
            elif kind == "blockquote_open":
                quote_depth += 1
            elif kind == "blockquote_close":
                quote_depth -= 1
            elif kind == "paragraph_open":
                inline = tokens[index + 1]
                children = inline.children or []
                if not lists and not quote_depth:
                    caption = _TABLE_CAPTION.match(inline.content.strip())
                    next_is_table = (
                        index + 3 < len(tokens) and tokens[index + 3].type == "table_open"
                    )
                    if caption and next_is_table:
                        pending_caption = caption.group(2)
                        index += 3
                        continue
                    if self._image_only(children):
                        self._figure(children)
                        index += 3
                        continue
                if lists:
                    context = lists[-1]
                    paragraph = self.document.add_paragraph(style="List Paragraph")
                    if item_started:
                        ooxml.set_numbering(paragraph, context.num_id, context.level)
                        item_started = False
                    else:
                        paragraph.paragraph_format.left_indent = Cm(0.635 * (context.level + 2))
                elif quote_depth:
                    paragraph = self.document.add_paragraph(style=LONG_QUOTE_STYLE)
                else:
                    paragraph = self.document.add_paragraph()
                self._inline(paragraph, children)
                index += 3
                continue
            elif kind in {"fence", "code_block"}:
                self._code(token.content)
            elif kind == "table_open":
                end = next(i for i in range(index, len(tokens)) if tokens[i].type == "table_close")
                self._table(tokens[index : end + 1], pending_caption)
                pending_caption = None
                index = end + 1
                continue
            elif kind == "hr":
                paragraph = self.document.add_paragraph()
                ooxml.paragraph_border(paragraph)
            elif kind == "html_block":
                content = token.content.strip().lower()
                if "pagebreak" in content:
                    self.page_break()
                elif re.match(r"<!--\s*toc\s*-->", content):
                    self.toc()
                elif not content.startswith("<!--"):
                    self.notes.append("Raw HTML blocks were skipped.")
            index += 1

    # -- pieces ----------------------------------------------------------------------

    def _inline(self, paragraph: Paragraph, children) -> None:
        bold = italic = strike = False
        link: str | None = None
        index = 0
        while index < len(children):
            child = children[index]
            kind = child.type
            if kind == "strong_open":
                bold = True
            elif kind == "strong_close":
                bold = False
            elif kind == "em_open":
                italic = True
            elif kind == "em_close":
                italic = False
            elif kind == "s_open":
                strike = True
            elif kind == "s_close":
                strike = False
            elif kind == "link_open":
                link = child.attrGet("href")
            elif kind == "link_close":
                link = None
            elif kind == "text":
                if link:
                    ooxml.add_hyperlink(paragraph, link, child.content)
                else:
                    run = paragraph.add_run(child.content)
                    run.bold, run.italic = bold or None, italic or None
                    run.font.strike = strike or None
            elif kind == "code_inline":
                run = paragraph.add_run(child.content, style=INLINE_CODE_STYLE)
                run.bold = bold or None
            elif kind == "softbreak":
                paragraph.add_run(" ")
            elif kind == "hardbreak":
                paragraph.add_run().add_break()
            elif kind == "html_inline":
                tag = child.content.lower()
                if tag.startswith("<br"):
                    paragraph.add_run().add_break()
            elif kind == "image":
                width, skip = self._width_attribute(children, index)
                self._picture(paragraph.add_run(), child.attrGet("src"), child.content, width)
                index += skip
            index += 1

    def _image_only(self, children) -> bool:
        meaningful = [
            c
            for c in children
            if not (c.type == "text" and (not c.content.strip() or _ATTRS.match(c.content.strip())))
            and c.type not in {"softbreak"}
        ]
        return len(meaningful) == 1 and meaningful[0].type == "image"

    def _width_attribute(self, children, index: int) -> tuple[float | None, int]:
        if index + 1 < len(children) and children[index + 1].type == "text":
            match = _ATTRS.match(children[index + 1].content.strip())
            if match:
                width = _width_cm(match.group(1), table_text_width_cm(self.document))
                rest = children[index + 1].content.strip()[match.end() :]
                children[index + 1].content = rest
                return width, 0 if rest else 1
        return None, 0

    def _figure(self, children) -> None:
        image = next(c for c in children if c.type == "image")
        position = children.index(image)
        width, _ = self._width_attribute(children, position)
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.keep_with_next = True
        placed = self._picture(paragraph.add_run(), image.attrGet("src"), image.content, width)
        if placed and image.content.strip():
            self._figures += 1
            self.caption(self.labels["figure"], image.content.strip())

    def _picture(self, run, src: str | None, alt: str, width_cm: float | None) -> bool:
        if not src or src.startswith(("http://", "https://", "data:")):
            self.notes.append(f"Image {src!r} skipped: only local files can be embedded.")
            return False
        try:
            path = self.resolve_image(src)
        except Exception as exc:  # noqa: BLE001 - reported, the rest of the document still builds
            self.notes.append(f"Image {src!r} skipped: {exc}")
            return False
        available = table_text_width_cm(self.document)
        if width_cm is None:
            width_cm = min(available, _natural_width_cm(path) or available)
        try:
            shape = run.add_picture(str(path), width=Cm(min(width_cm, available)))
        except Exception as exc:  # noqa: BLE001 - unsupported formats (svg, webp)
            self.notes.append(f"Image {src!r} skipped: {exc}")
            return False
        doc_pr = shape._inline.docPr
        doc_pr.set("descr", alt or Path(src).stem)
        return True

    def caption(self, label: str, text: str, *, above: bool = False) -> Paragraph:
        paragraph = self.document.add_paragraph(style="Caption")
        paragraph.add_run(f"{label} ")
        number = self._tables if label == self.labels["table"] else self._figures
        ooxml.add_field(paragraph, f"SEQ {label} \\* ARABIC", str(number))
        paragraph.add_run(f" – {text}")
        if above:
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(4)
        return paragraph

    def _code(self, content: str) -> None:
        lines = content.rstrip("\n").split("\n") or [""]
        for number, line in enumerate(lines):
            paragraph = self.document.add_paragraph(style=CODE_STYLE)
            paragraph.add_run(line.replace("\t", "    ") or " ")
            if number == len(lines) - 1:
                paragraph.paragraph_format.keep_with_next = False
                paragraph.paragraph_format.space_after = Pt(10)

    def _table(self, tokens, caption: str | None) -> None:
        rows: list[list[tuple[list, str]]] = []
        header_rows = 0
        in_head = False
        for position, token in enumerate(tokens):
            if token.type == "thead_open":
                in_head = True
            elif token.type == "thead_close":
                in_head = False
            elif token.type == "tr_open":
                rows.append([])
                if in_head:
                    header_rows += 1
            elif token.type in {"th_open", "td_open"}:
                style = token.attrGet("style") or ""
                align = style.split(":")[-1].strip() if "text-align" in style else ""
                inline = tokens[position + 1]
                rows[-1].append((inline.children or [], align))
        if not rows:
            return
        if caption:
            self._tables += 1
            self.caption(self.labels["table"], caption, above=True)
        columns = max(len(row) for row in rows)
        table = self.document.add_table(rows=len(rows), cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        style_table(table, self.preset, header_rows=header_rows)
        for r, row in enumerate(rows):
            for c in range(columns):
                cell = table.cell(r, c)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.line_spacing = 1.0
                if c < len(row):
                    children, align = row[c]
                    self._inline(paragraph, children)
                    if align in {"center", "right"}:
                        paragraph.alignment = (
                            WD_ALIGN_PARAGRAPH.CENTER
                            if align == "center"
                            else WD_ALIGN_PARAGRAPH.RIGHT
                        )
                    if r < header_rows:
                        for run in paragraph.runs:
                            run.bold = True
        spacer = self.document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)

    def page_break(self) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.add_run().add_break(WD_BREAK.PAGE)

    def toc(self, levels: int = 3) -> None:
        heading = self.document.add_paragraph(self.labels["toc"], style="TOC Heading")
        heading.paragraph_format.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if self.preset.name == "abnt" else WD_ALIGN_PARAGRAPH.LEFT
        )
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        ooxml.add_field(paragraph, f'TOC \\o "1-{levels}" \\h \\z \\u', self.labels["toc_hint"])
        ooxml.update_fields_on_open(self.document)
        self._toc_paragraph = paragraph
        self.page_break()

    def fill_toc_preview(self) -> None:
        """Put the headings in the TOC's cached result so viewers that do not update fields
        (LibreOffice previews, web viewers) still show a table of contents."""

        paragraph = getattr(self, "_toc_paragraph", None)
        if paragraph is None or not self.headings:
            return
        runs = paragraph._p.findall(qn("w:r"))
        placeholder = next(
            (r for r in runs if r.find(qn("w:t")) is not None and r.find(qn("w:fldChar")) is None),
            None,
        )
        if placeholder is None:
            return
        text = placeholder.find(qn("w:t"))
        text.text = ""
        entries = [(level, title) for level, title in self.headings if level <= 3]
        for number, (level, title) in enumerate(entries):
            if number:
                placeholder.append(placeholder.makeelement(qn("w:br"), {}))
            node = placeholder.makeelement(qn("w:t"), {qn("xml:space"): "preserve"})
            node.text = "    " * (level - 1) + title
            placeholder.append(node)
        placeholder.remove(text)


def style_table(table, preset: Preset, *, header_rows: int = 1) -> None:
    try:
        table.style = table.part.document.styles["Table Grid"]
    except KeyError:
        pass
    ooxml.table_borders(table, preset.table_border)
    for r, row in enumerate(table.rows):
        ooxml.keep_row_together(row)
        if r < header_rows:
            ooxml.repeat_header(row)
            for cell in row.cells:
                ooxml.shade(cell._tc.get_or_add_tcPr(), preset.table_header_fill)


def _width_cm(attributes: str, available: float) -> float | None:
    match = re.search(r"width\s*=\s*\"?([\d.]+)\s*(cm|mm|in|%)?", attributes)
    if not match:
        return None
    value, unit = float(match.group(1)), (match.group(2) or "cm")
    return {"cm": value, "mm": value / 10, "in": value * 2.54, "%": available * value / 100}[unit]


def _natural_width_cm(path: Path) -> float | None:
    try:
        from PIL import Image

        with Image.open(path) as image:
            dpi = image.info.get("dpi", (96, 96))[0] or 96
            return image.size[0] / float(dpi) * 2.54
    except Exception:  # noqa: BLE001
        return None


def fill_placeholders(text: str, title: str) -> list[tuple[str, str | None]]:
    """Split header/footer text into literal pieces and PAGE/NUMPAGES fields."""

    pieces: list[tuple[str, str | None]] = []
    text = text.replace("{title}", title).replace("{date}", date.today().strftime("%d/%m/%Y"))
    for part in re.split(r"(\{page\}|\{total\})", text):
        if part == "{page}":
            pieces.append(("1", "PAGE"))
        elif part == "{total}":
            pieces.append(("1", "NUMPAGES"))
        elif part:
            pieces.append((part, None))
    return pieces


def set_header_footer(
    document: Document,
    preset: Preset,
    *,
    header: str | None,
    footer: str | None,
    page_numbers: bool,
    title: str,
    skip_first_page: bool = False,
) -> None:
    position = preset.page_number_position
    if page_numbers and not any("{page}" in (t or "") for t in (header, footer)):
        if position.startswith("header"):
            header = f"{header}\t\t{{page}}" if header else "{page}"
        else:
            footer = f"{footer}\t\t{{page}}" if footer else "{page}"
    alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT if position.endswith("right") else WD_ALIGN_PARAGRAPH.CENTER
    )
    for section in document.sections:
        section.different_first_page_header_footer = skip_first_page
        for part, text in ((section.header, header), (section.footer, footer)):
            if text is None:
                continue
            part.is_linked_to_previous = False
            paragraph = part.paragraphs[0]
            for run in list(paragraph.runs):
                run._r.getparent().remove(run._r)
            paragraph.style = document.styles["Header" if part is section.header else "Footer"]
            if "\t" not in text:
                paragraph.alignment = alignment
            for literal, field_code in fill_placeholders(text, title):
                if field_code:
                    ooxml.add_field(paragraph, field_code, literal)
                else:
                    paragraph.add_run(literal)


def add_cover(document: Document, cover: dict, preset: Preset) -> None:
    """ABNT-style cover: institution on top, author, title in the middle, city and year below."""

    def line(text: str, *, bold=False, size=None, before=0.0, caps=True):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt = paragraph.paragraph_format
        fmt.first_line_indent = Cm(0)
        fmt.space_before = Pt(before)
        fmt.space_after = Pt(0)
        run = paragraph.add_run(text.upper() if caps else text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        return paragraph

    for item in str(cover.get("institution") or "").split("\n"):
        if item.strip():
            line(item.strip(), bold=True)
    if cover.get("author"):
        line(str(cover["author"]), before=90)
    line(str(cover.get("title") or ""), bold=True, size=preset.title_size, before=160)
    if cover.get("subtitle"):
        line(str(cover["subtitle"]), size=preset.body_size + 1, before=6, caps=False)
    if cover.get("description"):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(8)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(70)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = paragraph.add_run(str(cover["description"]))
        run.font.size = Pt(max(10.0, preset.body_size - 2))
    city_year = " ".join(str(cover.get(key) or "") for key in ("city", "year")).strip()
    if city_year:
        before = 200 if not cover.get("description") else 110
        line(
            city_year.replace(" ", "\n", 1)
            if cover.get("city") and cover.get("year")
            else city_year,
            before=before,
        )
    paragraph = document.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)
