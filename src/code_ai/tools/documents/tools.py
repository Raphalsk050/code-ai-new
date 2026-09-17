from __future__ import annotations

import asyncio
import os
import re
import tempfile
from pathlib import Path
from typing import Any

import docx
from docx.opc.exceptions import PackageNotFoundError

from code_ai.core.errors import ToolArgumentError, ToolExecutionError
from code_ai.tools.base import ToolCapability, ToolContext
from code_ai.tools.documents import builder, formatting, ooxml, operations, presets
from code_ai.tools.documents.inspect import inspect_document
from code_ai.tools.documents.to_markdown import document_to_markdown
from code_ai.tools.locations import LOCATION_SCHEMA, for_context
from code_ai.tools.office.common import (
    attach_images,
    clamp_int,
    display_path,
    optional_str,
    parse_page_spec,
    require_str,
    resolve_input,
    resolve_output,
)
from code_ai.tools.pdf.common import parse_json_argument
from code_ai.tools.schema import tool_schema

_DOCX = (".docx", ".docm", ".dotx")

_PRESET_FIELD = {
    "type": "string",
    "description": (
        "Look to apply: default (Calibri, blue headings), corporate (Arial, navy, compact), "
        "academic (Times 12pt, 1.5 spacing, justified, numbered headings), abnt (Brazilian "
        "ABNT NBR 14724: Arial 12, 1.5, margins 3/2/2/3 cm, numbered uppercase chapters, "
        "10pt long quotes indented 4 cm, page number top right), report (Cambria/Calibri, "
        "numbered), minimal (Segoe UI, airy)."
    ),
}
_RENDER_FIELD = {
    "type": "string",
    "description": (
        "Pages to render as images so you can check the layout, e.g. '1-2' (at most 12). "
        "Needs LibreOffice or Word; see install_converter."
    ),
}
_INSTALL_FIELD = {
    "type": "boolean",
    "description": (
        "When nothing can convert or render the document: install LibreOffice for this "
        "user (~350 MB download, no admin rights)."
    ),
}
_LANGUAGE_FIELD = {
    "type": "string",
    "description": "Document language, e.g. pt-BR or en-US; sets proofing and caption labels.",
}


def _open(path: Path):
    try:
        return docx.Document(str(path))
    except PackageNotFoundError:
        raise ToolArgumentError(f"{path.name} is not a valid .docx file.") from None
    except Exception as exc:  # noqa: BLE001 - corrupt packages raise many types
        raise ToolExecutionError(f"Could not open {path.name}: {exc}") from exc


def _save(document, target: Path) -> None:
    staging = target.with_name(f".{target.stem}.{os.getpid()}.tmp{target.suffix}")
    try:
        document.save(str(staging))
        os.replace(staging, target)
    finally:
        staging.unlink(missing_ok=True)


_PORTUGUESE = re.compile(r"\b(que|não|para|com|uma|são|dos|das|também|está)\b|ção\b|ções\b", re.I)
_ENGLISH = re.compile(r"\b(the|and|with|that|this|from|are|which|should)\b", re.I)


def _language(arguments: dict[str, Any], preset: presets.Preset, sample: str = "") -> str:
    given = optional_str(arguments, "language")
    if given:
        return given
    if preset.name == "abnt":
        return "pt-BR"
    portuguese = len(_PORTUGUESE.findall(sample))
    english = len(_ENGLISH.findall(sample))
    return "pt-BR" if portuguese > english else "en-US"


async def _preview(
    context: ToolContext, path: Path, arguments: dict[str, Any], payload: dict[str, Any]
) -> dict[str, Any]:
    spec = arguments.get("render_pages")
    if not spec:
        return payload
    from code_ai.tools.office.converters import convert
    from code_ai.tools.office.render import page_count, render_pages

    with tempfile.TemporaryDirectory(prefix="code-ai-doc-") as scratch:
        pdf = Path(scratch, path.stem + ".pdf")
        try:
            await convert(
                path,
                "pdf",
                pdf,
                verify_ssl=bool(context.config.ssl_verification),
                install=bool(arguments.get("install_converter")),
            )
        except ToolExecutionError as exc:
            payload["render_note"] = f"Saved, but not rendered: {exc}"
            return payload
        total = page_count(pdf)
        try:
            pages = parse_page_spec(spec, total)
        except ToolArgumentError:
            pages = list(range(min(total, 2)))
        payload["rendered_pages"] = [page + 1 for page in pages[:12]]
        payload["total_pages"] = total
        images = await asyncio.to_thread(render_pages, pdf, pages)
    return attach_images(payload, images)


def _image_resolver(context: ToolContext, location, base: Path):
    tree = for_context(context, location)

    def resolve(src: str) -> Path:
        candidate = (base / src) if not Path(src).is_absolute() else Path(src)
        return tree.resolve(candidate, must_exist=True)

    return resolve


def build_document(
    context: ToolContext, arguments: dict[str, Any], markdown: str, base: Path
) -> tuple[Any, list[str]]:
    preset = presets.with_overrides(
        presets.get_preset(optional_str(arguments, "preset", "default")),
        font=optional_str(arguments, "font") or None,
        size=arguments.get("font_size"),
    )
    language = _language(arguments, preset, markdown)
    template = optional_str(arguments, "template")
    if template:
        path = resolve_input(context, template, location=arguments.get("location"), suffixes=_DOCX)
        document = _open(path)
        body = document.element.body
        for child in list(body):
            if child.tag != ooxml.qn("w:sectPr"):
                body.remove(child)
        presets.ensure_styles(document)
    else:
        document = docx.Document()
        presets.apply_preset(document, preset)
        presets.apply_page_setup(
            document,
            preset,
            page_size=optional_str(arguments, "page_size") or None,
            orientation=optional_str(arguments, "orientation") or None,
            margins=presets.parse_margins(arguments.get("margins")),
        )
    ooxml.set_language(document, language)

    writer = builder.DocxWriter(
        document,
        preset,
        resolve_image=_image_resolver(context, arguments.get("location"), base),
        language=language,
    )
    number = arguments.get("number_headings")
    if (preset.heading_numbering if number is None else bool(number)) and not template:
        num_id = writer.numbering.heading_numbering(3)
        for level in range(3):
            ooxml.link_style_numbering(document.styles[f"Heading {level + 1}"], num_id, level)

    cover = parse_json_argument(arguments.get("cover"), "cover")
    if cover:
        if not isinstance(cover, dict):
            raise ToolArgumentError("'cover' is a JSON object: institution, author, title, ...")
        builder.add_cover(document, cover, preset)
    if arguments.get("toc") and "[toc]" not in markdown.lower() and "<!-- toc -->" not in markdown:
        writer.toc()
    writer.write_markdown(markdown)
    writer.fill_toc_preview()

    core = document.core_properties
    # python-docx's blank template claims to be written by "python-docx".
    core.author = optional_str(arguments, "author")
    core.last_modified_by = optional_str(arguments, "author")
    core.comments = ""
    title = optional_str(arguments, "title") or (cover or {}).get("title") or ""
    if not title and writer.headings is not None:
        first = next((p for p in document.paragraphs if p.style.name == "Title"), None)
        title = first.text if first is not None else ""
    core.title = title
    for key in ("author", "subject", "keywords"):
        value = optional_str(arguments, key)
        if value:
            setattr(core, key, value)
    header = arguments.get("header")
    footer = arguments.get("footer")
    page_numbers = arguments.get("page_numbers")
    page_numbers = True if page_numbers is None else bool(page_numbers)
    builder.set_header_footer(
        document,
        preset,
        header=str(header) if header else None,
        footer=str(footer) if footer else None,
        page_numbers=page_numbers,
        title=title,
        skip_first_page=bool(cover),
    )
    return document, writer.notes


class DocumentCreateTool:
    name = "document_create"
    description = (
        "Create a polished Word document (.docx) from Markdown. Supports headings (a lone "
        "leading '# ' becomes the title and '##' the first level), bold/italic/strike/`code`, "
        "real hyperlinks, nested bullet and numbered lists (each list restarts at 1), block "
        "quotes, code blocks, tables with alignment ('Table: caption' on the line before adds "
        "a numbered caption above), local images with '![caption](path){width=10cm}' "
        "(numbered figure captions, alt text), '\\newpage' page breaks and '[TOC]' for a table "
        "of contents. Styling goes through real Word styles from a preset, plus page size, "
        "margins, header/footer with {page} {total} {title} {date}, automatic heading "
        "numbering, an ABNT-style cover page, document properties and language. Or base it "
        "on a template .docx to reuse its styles. Use render_pages to see the result."
    )
    capabilities = frozenset(
        {ToolCapability.LOCAL_READ, ToolCapability.LOCAL_WRITE, ToolCapability.PROCESS}
    )
    input_schema = tool_schema(
        {
            "output": {"type": "string", "description": "Path of the .docx to write."},
            "content": {"type": "string", "description": "The document body as Markdown."},
            "source": {
                "type": "string",
                "description": "A Markdown file to use instead of content; image paths are "
                "relative to it.",
            },
            "location": LOCATION_SCHEMA,
            "preset": _PRESET_FIELD,
            "template": {
                "type": "string",
                "description": "Existing .docx/.dotx whose styles and page setup to reuse.",
            },
            "title": {"type": "string", "description": "Document title property."},
            "author": {"type": "string", "description": "Author property."},
            "subject": {"type": "string", "description": "Subject property."},
            "keywords": {"type": "string", "description": "Keywords property."},
            "language": _LANGUAGE_FIELD,
            "page_size": {"type": "string", "description": "a4 (default), letter, legal, a5, a3."},
            "orientation": {"type": "string", "description": "portrait (default) or landscape."},
            "margins": {
                "type": "string",
                "description": "Margins 'top right bottom left', e.g. '3cm 2cm 2cm 3cm'.",
            },
            "font": {"type": "string", "description": "Override the preset's font family."},
            "font_size": {"type": "number", "description": "Override the body font size (pt)."},
            "header": {"type": "string", "description": "Header text; may use {page} {total}."},
            "footer": {"type": "string", "description": "Footer text; may use {page} {total}."},
            "page_numbers": {
                "type": "boolean",
                "description": "Add page numbers where the preset puts them (default true).",
            },
            "toc": {"type": "boolean", "description": "Add a table of contents at the start."},
            "number_headings": {
                "type": "boolean",
                "description": "Number headings 1, 1.1, 1.1.1 (default from the preset).",
            },
            "cover": {
                "type": "string",
                "description": (
                    'JSON cover page, e.g. {"institution":"Universidade X\\nFaculdade Y",'
                    '"author":"Nome","title":"Título","subtitle":"...","description":"Trabalho '
                    'apresentado a...","city":"Manaus","year":"2026"}.'
                ),
            },
            "overwrite": {
                "type": "boolean",
                "description": "Replace output if it exists (default true).",
            },
            "render_pages": _RENDER_FIELD,
            "install_converter": _INSTALL_FIELD,
        },
        required=("output",),
    )

    async def execute(self, arguments: dict[str, Any], context: ToolContext) -> dict[str, Any]:
        location = arguments.get("location")
        output = resolve_output(
            context,
            require_str(arguments, "output", self.name),
            location=location,
            overwrite=arguments.get("overwrite") is not False,
            suffixes=(".docx",),
        )
        content = arguments.get("content")
        source = optional_str(arguments, "source")
        if isinstance(content, str) and content.strip():
            markdown, base = content, output.parent
        elif source:
            path = resolve_input(context, source, location=location)
            markdown, base = path.read_text(encoding="utf-8", errors="replace"), path.parent
        else:
            raise ToolArgumentError("document_create needs 'content' (Markdown) or 'source'.")

        def work():
            document, notes = build_document(context, arguments, markdown, base)
            _save(document, output)
            return document, notes

        document, notes = await asyncio.to_thread(work)
        payload: dict[str, Any] = {
            "path": display_path(context, output),
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
            "images": len(document.inline_shapes),
            "preset": optional_str(arguments, "preset", "default"),
        }
        if notes:
            payload["notes"] = sorted(set(notes))
        return await _preview(context, output, arguments, payload)


class DocumentInspectTool:
    name = "document_inspect"
    description = (
        "Read a Word document's structure and report its formatting problems. Covers properties, "
        "page setup, headers/footers, outline "
        "of headings, every paragraph with its index (p0, p1... - the numbers document_edit "
        "uses), style and list membership, tables (t0, t1...), styles in use with fonts, "
        "tracked changes and comments. Also reports formatting problems: direct formatting "
        "fighting the styles, mixed body fonts, empty paragraphs used as spacing, list numbers "
        "typed by hand, bold lines posing as headings, skipped heading levels, images without "
        "alt text, long tables without a repeating header. Optionally renders pages."
    )
    capabilities = frozenset({ToolCapability.LOCAL_READ, ToolCapability.PROCESS})
    input_schema = tool_schema(
        {
            "path": {"type": "string", "description": "The .docx to inspect."},
            "location": LOCATION_SCHEMA,
            "offset": {
                "type": "integer",
                "description": "First paragraph index to list (default 0).",
            },
            "limit": {
                "type": "integer",
                "description": "How many paragraphs to list (default 200).",
            },
            "text_chars": {
                "type": "integer",
                "description": "Characters of text shown per paragraph (default 160).",
            },
            "render_pages": _RENDER_FIELD,
            "install_converter": _INSTALL_FIELD,
        },
        required=("path",),
    )

    async def execute(self, arguments: dict[str, Any], context: ToolContext) -> dict[str, Any]:
        path = resolve_input(
            context,
            require_str(arguments, "path", self.name),
            location=arguments.get("location"),
            suffixes=_DOCX,
        )

        def work():
            document = _open(path)
            return inspect_document(
                document,
                offset=clamp_int(arguments.get("offset"), default=0, low=0, high=10**6),
                limit=clamp_int(arguments.get("limit"), default=200, low=1, high=2000),
                text_chars=clamp_int(arguments.get("text_chars"), default=160, low=20, high=5000),
            )

        report = await asyncio.to_thread(work)
        payload = {"path": display_path(context, path), **report}
        return await _preview(context, path, arguments, payload)


_EDIT_DOC = (
    "JSON array of operations, applied in order; nothing is saved unless all succeed. "
    "Paragraph indexes are the p numbers from document_inspect (0-based; '3-7,9' ranges) "
    "and are re-read after each operation. Ops: "
    "replace {find, replace, regex?, match_case? true, whole_word?, count?, paragraphs?, "
    "scope? 'all'|'body'} (works across formatting runs, tables, headers) | "
    "format_text {find, regex?, bold?, italic?, underline?, strike?, color? '#c00000', "
    "highlight? yellow|green|cyan|pink|gray|red|none, font?, size?} | "
    "insert {markdown, after? p | before? p | after_heading? 'text' (end of that section) | "
    "after_table? t | at? 'start'|'end'} (Markdown as in document_create) | "
    "delete {paragraphs | section: 'heading text' (whole section) | table} | "
    "set_style {paragraphs, style} | paragraph_format {paragraphs, alignment?, space_before?, "
    "space_after?, line_spacing?, first_line_indent? cm, left_indent? cm, keep_with_next?, "
    "page_break_before?} | table_cell {table, row, column, text, bold?} | table_add_row "
    "{table, values: [...], after?} | table_delete_row {table, row} | page_break {after | "
    "before} | toc {after? | before? | at?, levels?} | header_footer {header?, footer?, "
    "page_numbers?, position? 'footer-center'|'footer-right'|'header-right', "
    "skip_first_page?} | properties {title?, author?, subject?, keywords?, comments?} | "
    "page_setup {page_size?, orientation?, margins?} | comment {paragraph, text, find?, "
    "author?} | accept_changes {} | reject_changes {}. "
    'Example: [{"op":"replace","find":"ACME","replace":"Contoso"},'
    '{"op":"insert","after_heading":"Introdução","markdown":"Novo parágrafo com **ênfase**."}]'
)


class DocumentEditTool:
    name = "document_edit"
    description = (
        "Edit an existing Word document (.docx) with a list of operations. Saved atomically: find "
        "and replace (even across formatting runs, in tables and headers), format matching "
        "text, insert Markdown content after a paragraph or at the end of a named section, "
        "delete paragraphs, sections or tables, set styles and paragraph formatting, edit "
        "table cells and rows, page breaks, table of contents, headers/footers and page "
        "numbers, properties, page setup, comments, and accepting or rejecting tracked "
        "changes. Run document_inspect first for paragraph indexes."
    )
    capabilities = frozenset(
        {ToolCapability.LOCAL_READ, ToolCapability.LOCAL_WRITE, ToolCapability.PROCESS}
    )
    input_schema = tool_schema(
        {
            "path": {"type": "string", "description": "The .docx to edit."},
            "location": LOCATION_SCHEMA,
            "operations": {"type": "string", "description": _EDIT_DOC},
            "output": {"type": "string", "description": "Write here instead of in place."},
            "language": _LANGUAGE_FIELD,
            "render_pages": _RENDER_FIELD,
            "install_converter": _INSTALL_FIELD,
        },
        required=("path", "operations"),
    )

    async def execute(self, arguments: dict[str, Any], context: ToolContext) -> dict[str, Any]:
        location = arguments.get("location")
        source = resolve_input(
            context, require_str(arguments, "path", self.name), location=location, suffixes=_DOCX
        )
        specs = operations.normalize(parse_json_argument(arguments.get("operations"), "operations"))
        output_raw = optional_str(arguments, "output")
        target = (
            resolve_output(context, output_raw, location=location, suffixes=(".docx",))
            if output_raw
            else source
        )

        def work():
            document = _open(source)
            sample = " ".join(p.text for p in document.paragraphs[:80])
            language = _language(arguments, presets.PRESETS["default"], sample)
            state = operations.DocState(
                document=document,
                resolve_image=_image_resolver(context, location, source.parent),
                language=language,
            )
            done = operations.apply(state, specs)
            _save(document, target)
            return document, done, state.notes

        document, done, notes = await asyncio.to_thread(work)
        payload: dict[str, Any] = {
            "path": display_path(context, target),
            "applied": done,
            "paragraphs": len(document.paragraphs),
        }
        if notes:
            payload["notes"] = notes
        return await _preview(context, target, arguments, payload)


class DocumentFormatTool:
    name = "document_format"
    description = (
        "Reformat a Word document (.docx) consistently onto a style preset such as ABNT. "
        "It restyles "
        "Normal, headings, title, captions, quotes, lists and tables; applies page size and "
        "margins; strips direct font/size/colour and spacing that fight the styles (keeping "
        "bold, italic and underline); removes empty spacer paragraphs; turns bold lines that "
        "act as headings into real headings and hand-typed '1.' or '-' items into real "
        "lists; gives 'Figura 1...' lines the Caption style; centres images; numbers headings; "
        "adds page numbers and optionally a table of contents. Reports every change by count. "
        "Switch individual fixes off with their flags."
    )
    capabilities = frozenset(
        {ToolCapability.LOCAL_READ, ToolCapability.LOCAL_WRITE, ToolCapability.PROCESS}
    )
    input_schema = tool_schema(
        {
            "path": {"type": "string", "description": "The .docx to format."},
            "location": LOCATION_SCHEMA,
            "preset": _PRESET_FIELD,
            "output": {"type": "string", "description": "Write here instead of in place."},
            "font": {"type": "string", "description": "Override the preset's font family."},
            "font_size": {"type": "number", "description": "Override the body size (pt)."},
            "page_size": {"type": "string", "description": "a4, letter, legal, a5, a3."},
            "orientation": {"type": "string", "description": "portrait or landscape."},
            "margins": {
                "type": "string",
                "description": "'top right bottom left', e.g. '3cm 2cm 2cm 3cm'.",
            },
            "keep_page_setup": {
                "type": "boolean",
                "description": "Leave page size and margins alone.",
            },
            "clear_direct_formatting": {"type": "boolean", "description": "Default true."},
            "remove_empty_paragraphs": {"type": "boolean", "description": "Default true."},
            "fix_fake_headings": {"type": "boolean", "description": "Default true."},
            "fix_manual_lists": {"type": "boolean", "description": "Default true."},
            "style_tables": {"type": "boolean", "description": "Default true."},
            "page_numbers": {"type": "boolean", "description": "Default true."},
            "toc": {
                "type": "boolean",
                "description": "Insert a table of contents before the first heading.",
            },
            "number_headings": {"type": "boolean", "description": "Default from the preset."},
            "language": _LANGUAGE_FIELD,
            "render_pages": _RENDER_FIELD,
            "install_converter": _INSTALL_FIELD,
        },
        required=("path",),
    )

    async def execute(self, arguments: dict[str, Any], context: ToolContext) -> dict[str, Any]:
        location = arguments.get("location")
        source = resolve_input(
            context, require_str(arguments, "path", self.name), location=location, suffixes=_DOCX
        )
        output_raw = optional_str(arguments, "output")
        target = (
            resolve_output(context, output_raw, location=location, suffixes=(".docx",))
            if output_raw
            else source
        )
        preset = presets.with_overrides(
            presets.get_preset(optional_str(arguments, "preset", "default")),
            font=optional_str(arguments, "font") or None,
            size=arguments.get("font_size"),
        )

        def flag(name: str, default: bool) -> bool:
            value = arguments.get(name)
            return default if value is None else bool(value)

        options = formatting.FormatOptions(
            clear_direct_formatting=flag("clear_direct_formatting", True),
            remove_empty_paragraphs=flag("remove_empty_paragraphs", True),
            fix_fake_headings=flag("fix_fake_headings", True),
            fix_manual_lists=flag("fix_manual_lists", True),
            style_tables=flag("style_tables", True),
            page_numbers=flag("page_numbers", True),
            toc=flag("toc", False),
            number_headings=arguments.get("number_headings"),
            page_size=optional_str(arguments, "page_size") or None,
            orientation=optional_str(arguments, "orientation") or None,
            margins=presets.parse_margins(arguments.get("margins")),
            keep_page_setup=flag("keep_page_setup", False),
        )

        def work():
            document = _open(source)
            sample = " ".join(p.text for p in document.paragraphs[:80])
            language = _language(arguments, preset, sample)
            ooxml.set_language(document, language)
            changes = formatting.format_document(document, preset, options, language)
            _save(document, target)
            return changes

        changes = await asyncio.to_thread(work)
        payload = {"path": display_path(context, target), "preset": preset.name, "changes": changes}
        return await _preview(context, target, arguments, payload)


class DocumentConvertTool:
    name = "document_convert"
    description = (
        "Convert documents; formats come from the file extensions. Markdown (.md) to .docx "
        "(same engine and options as document_create: preset, toc...). .docx to Markdown "
        "(headings, lists, tables, emphasis, links; images extracted next to it). Between "
        ".docx, .doc, .odt, .rtf, .html, .txt and .pdf through LibreOffice or Microsoft Word "
        "(install_converter installs LibreOffice when neither exists). render_pages shows "
        "the pages of a produced PDF or document."
    )
    capabilities = frozenset(
        {
            ToolCapability.LOCAL_READ,
            ToolCapability.LOCAL_WRITE,
            ToolCapability.PROCESS,
            ToolCapability.WEB,
        }
    )
    input_schema = tool_schema(
        {
            "source": {"type": "string", "description": "File to convert."},
            "output": {
                "type": "string",
                "description": "File to write; its extension picks the format.",
            },
            "location": LOCATION_SCHEMA,
            "preset": _PRESET_FIELD,
            "toc": {
                "type": "boolean",
                "description": "Markdown to .docx: add a table of contents.",
            },
            "language": _LANGUAGE_FIELD,
            "overwrite": {
                "type": "boolean",
                "description": "Replace output if it exists (default true).",
            },
            "render_pages": _RENDER_FIELD,
            "install_converter": _INSTALL_FIELD,
        },
        required=("source", "output"),
    )

    async def execute(self, arguments: dict[str, Any], context: ToolContext) -> dict[str, Any]:
        location = arguments.get("location")
        source = resolve_input(
            context, require_str(arguments, "source", self.name), location=location
        )
        output = resolve_output(
            context,
            require_str(arguments, "output", self.name),
            location=location,
            overwrite=arguments.get("overwrite") is not False,
        )
        src, dst = source.suffix.lower(), output.suffix.lower()
        if src == dst:
            raise ToolArgumentError("source and output have the same format.")

        if src in {".md", ".markdown", ".txt"} and dst == ".docx":
            markdown = source.read_text(encoding="utf-8", errors="replace")

            def build():
                document, notes = build_document(context, arguments, markdown, source.parent)
                _save(document, output)
                return notes

            notes = await asyncio.to_thread(build)
            payload: dict[str, Any] = {"path": display_path(context, output), "via": "markdown"}
            if notes:
                payload["notes"] = sorted(set(notes))
            return await _preview(context, output, arguments, payload)

        if src in _DOCX and dst in {".md", ".markdown"}:
            media_dir = output.parent / f"{output.stem}_media"

            def export():
                document = _open(source)
                markdown, images = document_to_markdown(document, media_dir, media_dir.name)
                output.write_text(markdown, encoding="utf-8")
                return images, len(markdown)

            images, chars = await asyncio.to_thread(export)
            payload = {"path": display_path(context, output), "chars": chars, "images": images}
            if images:
                payload["media"] = display_path(context, media_dir)
            return payload

        from code_ai.tools.office.converters import convert

        conversion = await convert(
            source,
            dst.lstrip("."),
            output,
            verify_ssl=bool(context.config.ssl_verification),
            install=bool(arguments.get("install_converter")),
        )
        payload = {"path": display_path(context, output), "via": conversion.backend}
        if dst == ".pdf" and arguments.get("render_pages"):
            from code_ai.tools.office.render import page_count, render_pages

            total = page_count(output)
            try:
                pages = parse_page_spec(arguments["render_pages"], total)
            except ToolArgumentError:
                pages = list(range(min(total, 2)))
            payload["total_pages"] = total
            return attach_images(payload, await asyncio.to_thread(render_pages, output, pages))
        if dst in _DOCX:
            return await _preview(context, output, arguments, payload)
        return payload
