"""Normalising an existing document onto a preset, reporting every change it makes."""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass

from docx.document import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

from code_ai.tools.documents import builder, ooxml, presets
from code_ai.tools.documents.inspect import MANUAL_NUMBERING, heading_level

_KEEP_RUN_PROPS = {"w:b", "w:bCs", "w:i", "w:iCs", "w:u", "w:strike", "w:vertAlign", "w:rStyle"}
_DROP_RUN_PROPS = ("w:rFonts", "w:sz", "w:szCs", "w:color", "w:highlight", "w:shd", "w:caps")
_DROP_PARAGRAPH_PROPS = ("w:jc", "w:spacing", "w:ind")
_NUMBERED_HEADING = re.compile(r"^\s*(\d+(?:\.\d+)*)\.?\s+(?=\S)")
_CAPTION = re.compile(r"^\s*(figura|figure|tabela|table|quadro|gráfico|grafico)\s+\d+", re.I)
_BULLET_PREFIX = re.compile(r"^\s*[-•▪◦*]\s+")
_NUMBER_PREFIX = re.compile(r"^\s*(\d{1,3}|[a-zA-Z]|[ivxIVX]{1,4})[.)]\s+")


@dataclass
class FormatOptions:
    clear_direct_formatting: bool = True
    remove_empty_paragraphs: bool = True
    fix_fake_headings: bool = True
    fix_manual_lists: bool = True
    style_tables: bool = True
    page_numbers: bool = True
    toc: bool = False
    number_headings: bool | None = None
    page_size: str | None = None
    orientation: str | None = None
    margins: tuple[float, float, float, float] | None = None
    keep_page_setup: bool = False


def format_document(
    document: Document, preset: presets.Preset, options: FormatOptions, language: str
) -> dict:
    changes: Counter[str] = Counter()
    presets.ensure_styles(document)
    presets.apply_preset(document, preset)
    changes["styles restyled to the preset"] = 1
    if not options.keep_page_setup:
        presets.apply_page_setup(
            document,
            preset,
            page_size=options.page_size,
            orientation=options.orientation,
            margins=options.margins,
        )
        changes["page setup applied"] = len(document.sections)

    if options.fix_fake_headings:
        changes["fake headings turned into real headings"] = _fix_fake_headings(document)
    if options.fix_manual_lists:
        changes["typed list items turned into real lists"] = _fix_manual_lists(document)
    if options.clear_direct_formatting:
        runs, paragraphs = _clear_direct_formatting(document)
        changes["runs stripped of direct font/size/colour"] = runs
        changes["paragraphs stripped of direct alignment/spacing/indent"] = paragraphs
    if options.remove_empty_paragraphs:
        changes["empty spacer paragraphs removed"] = _remove_empty_paragraphs(document)
    changes["captions given the Caption style"] = _style_captions(document)
    changes["image paragraphs centred"] = _center_images(document)
    if options.style_tables:
        for table in document.tables:
            builder.style_table(table, preset, header_rows=1)
        changes["tables styled"] = len(document.tables)

    number = (
        preset.heading_numbering if options.number_headings is None else options.number_headings
    )
    if number:
        stripped = _strip_typed_heading_numbers(document)
        numbering = ooxml.Numbering(document)
        num_id = numbering.heading_numbering(3)
        for level in range(3):
            ooxml.link_style_numbering(document.styles[f"Heading {level + 1}"], num_id, level)
        changes["headings numbered automatically"] = 1
        changes["typed heading numbers removed"] = stripped

    if options.toc:
        _insert_toc(document, preset, language)
        changes["table of contents inserted"] = 1
    if options.page_numbers:
        builder.set_header_footer(
            document,
            preset,
            header=None,
            footer=None,
            page_numbers=True,
            title=document.core_properties.title or "",
        )
        changes["page numbers added"] = 1
    ooxml.update_fields_on_open(document)
    return {key: value for key, value in changes.items() if value}


def _is_body(paragraph: Paragraph) -> bool:
    name = paragraph.style.name if paragraph.style is not None else "Normal"
    return name in {"Normal", "Body Text", "Body Text 2", "Body Text 3", "No Spacing", "Default"}


def _fix_fake_headings(document: Document) -> int:
    paragraphs = document.paragraphs
    fixed = 0
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        runs = [run for run in paragraph.runs if run.text.strip()]
        if (
            not _is_body(paragraph)
            or not runs
            or len(text) > 90
            or text.endswith((".", ":", ";", ","))
            or not all(run.bold for run in runs)
            or index + 1 >= len(paragraphs)
            or not paragraphs[index + 1].text.strip()
            or ooxml.list_info(paragraph) is not None
        ):
            continue
        letters = [ch for ch in text if ch.isalpha()]
        upper = letters and sum(ch.isupper() for ch in letters) / len(letters) > 0.8
        numbered = _NUMBERED_HEADING.match(text)
        depth = numbered.group(1).count(".") + 1 if numbered else (1 if upper else 2)
        paragraph.style = document.styles[f"Heading {min(depth, 3)}"]
        for run in runs:
            run.bold = None
        fixed += 1
    return fixed


def _fix_manual_lists(document: Document) -> int:
    numbering = ooxml.Numbering(document)
    converted = 0
    current_kind: str | None = None
    num_id = 0
    for paragraph in document.paragraphs:
        text = paragraph.text
        if (
            not _is_body(paragraph)
            or ooxml.list_info(paragraph) is not None
            or not MANUAL_NUMBERING.match(text)
        ):
            current_kind = None
            continue
        kind = "bullet" if _BULLET_PREFIX.match(text) else "number"
        if kind != current_kind:
            num_id = numbering.bullet_list() if kind == "bullet" else numbering.ordered_list()
            current_kind = kind
        prefix = (_BULLET_PREFIX if kind == "bullet" else _NUMBER_PREFIX).match(text)
        if prefix is None:
            current_kind = None
            continue
        _remove_prefix(paragraph, len(prefix.group(0)))
        paragraph.style = document.styles["List Paragraph"]
        ooxml.set_numbering(paragraph, num_id, 0)
        converted += 1
    return converted


def _remove_prefix(paragraph: Paragraph, length: int) -> None:
    remaining = length
    for run in paragraph.runs:
        if remaining <= 0:
            break
        text = run.text
        cut = min(len(text), remaining)
        run.text = text[cut:]
        remaining -= cut


def _clear_direct_formatting(document: Document) -> tuple[int, int]:
    runs_cleared = paragraphs_cleared = 0
    targets = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                targets.extend(cell.paragraphs)
    for paragraph in targets:
        name = paragraph.style.name if paragraph.style is not None else ""
        if name in {presets.CODE_STYLE, "Title", "Subtitle"}:
            continue
        for run in paragraph.runs:
            r_pr = run._r.rPr
            if r_pr is None:
                continue
            style = r_pr.find(qn("w:rStyle"))
            if style is not None and style.get(qn("w:val")) in {
                "InlineCode",
                presets.INLINE_CODE_STYLE,
            }:
                continue
            removed = False
            for tag in _DROP_RUN_PROPS:
                for element in r_pr.findall(qn(tag)):
                    r_pr.remove(element)
                    removed = True
            runs_cleared += removed
        in_table = paragraph._p.getparent() is not None and paragraph._p.getparent().tag == qn(
            "w:tc"
        )
        p_pr = paragraph._p.pPr
        if p_pr is None or in_table:
            continue
        has_image = paragraph._p.find(f".//{qn('w:drawing')}") is not None
        removed = False
        for tag in _DROP_PARAGRAPH_PROPS:
            if tag == "w:ind" and p_pr.find(qn("w:numPr")) is not None:
                continue
            if tag == "w:jc" and has_image:
                continue
            for element in p_pr.findall(qn(tag)):
                p_pr.remove(element)
                removed = True
        paragraphs_cleared += removed
    return runs_cleared, paragraphs_cleared


def _remove_empty_paragraphs(document: Document) -> int:
    removed = 0
    for paragraph in list(document.paragraphs):
        element = paragraph._p
        if paragraph.text.strip():
            continue
        if any(
            element.find(f".//{qn(tag)}") is not None
            for tag in ("w:drawing", "w:br", "w:sectPr", "w:fldChar", "w:object", "w:pict")
        ):
            continue
        element.getparent().remove(element)
        removed += 1
    return removed


def _style_captions(document: Document) -> int:
    styled = 0
    for paragraph in document.paragraphs:
        if _is_body(paragraph) and _CAPTION.match(paragraph.text) and len(paragraph.text) < 200:
            paragraph.style = document.styles["Caption"]
            styled += 1
    return styled


def _center_images(document: Document) -> int:
    centred = 0
    for paragraph in document.paragraphs:
        if paragraph.text.strip() or paragraph._p.find(f".//{qn('w:drawing')}") is None:
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        centred += 1
    return centred


def _strip_typed_heading_numbers(document: Document) -> int:
    stripped = 0
    for paragraph in document.paragraphs:
        level = heading_level(paragraph)
        if not level:
            continue
        match = _NUMBERED_HEADING.match(paragraph.text)
        if match:
            _remove_prefix(paragraph, len(match.group(0)))
            stripped += 1
    return stripped


def _insert_toc(document: Document, preset: presets.Preset, language: str) -> None:
    body = document.element.body
    first_heading = next(
        (p for p in document.paragraphs if heading_level(p) not in (None, 0)), None
    )
    content = [child for child in body if child.tag != qn("w:sectPr")]
    before = len(content)
    writer = builder.DocxWriter(document, preset, resolve_image=lambda s: s, language=language)
    writer.toc()
    created = [child for child in body if child.tag != qn("w:sectPr")][before:]
    if first_heading is None:
        return
    for element in created:
        first_heading._p.addprevious(element)
