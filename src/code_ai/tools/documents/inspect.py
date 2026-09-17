"""The shape of a .docx and what is wrong with its formatting."""

from __future__ import annotations

import re
from collections import Counter
from typing import Any

from docx.document import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from code_ai.tools.documents.ooxml import list_info

MANUAL_NUMBERING = re.compile(r"^\s*(\d{1,3}[.)]|[a-zA-Z][.)]|[ivxIVX]{1,4}[.)]|[-•▪◦*])\s+\S")
MAX_ISSUE_EXAMPLES = 8


def heading_level(paragraph: Paragraph) -> int | None:
    name = paragraph.style.name if paragraph.style is not None else ""
    if name == "Title":
        return 0
    match = re.match(r"Heading (\d)", name)
    if match:
        return int(match.group(1))
    p_pr = paragraph._p.pPr
    outline = p_pr.find(qn("w:outlineLvl")) if p_pr is not None else None
    if outline is not None:
        level = int(outline.get(qn("w:val")))
        return level + 1 if level < 9 else None
    return None


def blocks(document: Document):
    """Body paragraphs and tables in reading order, each with its own running index."""

    paragraph_index = table_index = 0
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            yield "p", paragraph_index, item
            paragraph_index += 1
        elif isinstance(item, Table):
            yield "t", table_index, item
            table_index += 1


def inspect_document(document: Document, *, offset: int, limit: int, text_chars: int) -> dict:
    core = document.core_properties
    report: dict[str, Any] = {
        "properties": {
            key: value
            for key, value in {
                "title": core.title,
                "author": core.author,
                "subject": core.subject,
                "keywords": core.keywords,
                "last_modified_by": core.last_modified_by,
            }.items()
            if value
        },
        "sections": [_section(section) for section in document.sections[:10]],
        "counts": {
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
            "images": len(document.inline_shapes),
            "words": sum(len(p.text.split()) for p in document.paragraphs),
        },
    }
    outline = []
    for kind, index, paragraph in blocks(document):
        level = heading_level(paragraph) if kind == "p" else None
        if level is not None and paragraph.text.strip():
            indent = "  " * max(0, level - 1)
            outline.append(f"{indent}[p{index}] {paragraph.text.strip()[:90]}")
    report["outline"] = outline[:120]

    listing = []
    last_paragraph = -1
    for kind, index, item in blocks(document):
        if kind == "p":
            last_paragraph = index
            if index < offset or index >= offset + limit:
                continue
            entry = f"p{index} [{item.style.name if item.style is not None else '?'}]"
            info = list_info(item)
            if info:
                entry += f" ({info[0]} list, level {info[1]})"
            text = item.text.strip()
            if not text and item._p.find(f".//{qn('w:drawing')}") is not None:
                text = "<image>"
            listing.append(f"{entry}: {text[:text_chars]}{'...' if len(text) > text_chars else ''}")
        else:
            if not offset - 1 <= last_paragraph < offset + limit:
                continue
            rows = [" | ".join(c.text.strip()[:25] for c in row.cells) for row in item.rows[:3]]
            listing.append(
                f"t{index} table {len(item.rows)}x{len(item.columns)}: " + " / ".join(rows)
            )
    report["blocks"] = listing
    if offset + limit < len(document.paragraphs):
        report["more"] = f"paragraphs {offset + limit}+ not listed; pass offset to page on"

    report["styles_in_use"] = _styles_in_use(document)
    report["review"] = {
        "tracked_changes": len(document.element.body.findall(f".//{qn('w:ins')}"))
        + len(document.element.body.findall(f".//{qn('w:del')}")),
        "comments": _comment_count(document),
    }
    report["issues"] = find_issues(document)
    return report


def _section(section) -> dict[str, Any]:
    def cm(value):
        return round(value / 360000, 2) if value is not None else None

    return {
        "page_cm": [cm(section.page_width), cm(section.page_height)],
        "orientation": "landscape" if section.page_width > section.page_height else "portrait",
        "margins_cm": {
            "top": cm(section.top_margin),
            "right": cm(section.right_margin),
            "bottom": cm(section.bottom_margin),
            "left": cm(section.left_margin),
        },
        "header": " ".join(p.text for p in section.header.paragraphs).strip()[:120],
        "footer": " ".join(p.text for p in section.footer.paragraphs).strip()[:120],
    }


def _comment_count(document: Document) -> int:
    try:
        return len(list(document.comments))
    except Exception:  # noqa: BLE001 - older files without a comments part
        return 0


def _effective_font(paragraph: Paragraph, run) -> tuple[str | None, float | None]:
    name = run.font.name
    size = run.font.size.pt if run.font.size else None
    style = paragraph.style
    while style is not None and (name is None or size is None):
        name = name or style.font.name
        if size is None and style.font.size is not None:
            size = style.font.size.pt
        style = style.base_style
    return name, size


def _styles_in_use(document: Document) -> list[str]:
    counter: Counter[str] = Counter()
    for paragraph in document.paragraphs:
        counter[paragraph.style.name if paragraph.style is not None else "?"] += 1
    described = []
    for name, count in counter.most_common(25):
        style = document.styles[name] if name in [s.name for s in document.styles] else None
        font = ""
        if style is not None and style.font is not None:
            size = f" {style.font.size.pt:g}pt" if style.font.size else ""
            font = f" ({style.font.name or 'inherited'}{size})"
        described.append(f"{name}{font}: {count}")
    return described


def find_issues(document: Document) -> list[dict[str, Any]]:
    issues: dict[str, dict[str, Any]] = {}

    def flag(rule: str, message: str, where: str) -> None:
        entry = issues.setdefault(rule, {"rule": rule, "message": message, "count": 0, "at": []})
        entry["count"] += 1
        if len(entry["at"]) < MAX_ISSUE_EXAMPLES:
            entry["at"].append(where)

    body_fonts: Counter[tuple] = Counter()
    empty_run = 0
    previous_level = 0
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style is not None else ""
        level = heading_level(paragraph)

        if not text and paragraph._p.find(f".//{qn('w:drawing')}") is None:
            if paragraph._p.find(f".//{qn('w:br')}") is None:
                empty_run += 1
                if empty_run == 2:
                    flag("empty-paragraphs", "Empty paragraphs used as spacing", f"p{index}")
            continue
        empty_run = 0

        if level:
            if previous_level and level > previous_level + 1:
                flag(
                    "heading-skip",
                    "Heading levels skip (e.g. Heading 1 straight to Heading 3)",
                    f"p{index}",
                )
            previous_level = level
            continue

        runs = [run for run in paragraph.runs if run.text.strip()]
        if style_name in {"Normal", "Body Text"}:
            for run in runs:
                body_fonts[_effective_font(paragraph, run)] += len(run.text)
            direct = [r for r in runs if r.font.name or r.font.size or r.font.color.rgb]
            if direct and len(direct) == len(runs):
                flag(
                    "direct-formatting",
                    "Font, size or colour set directly instead of through a style",
                    f"p{index}",
                )
            if (
                runs
                and len(text) < 90
                and not text.endswith((".", ":", ";", ","))
                and all(run.bold for run in runs)
                and index + 1 < len(paragraphs)
                and paragraphs[index + 1].text.strip()
            ):
                flag("fake-heading", "Bold short paragraph acting as a heading", f"p{index}")
            if MANUAL_NUMBERING.match(text) and list_info(paragraph) is None:
                flag(
                    "manual-numbering",
                    "List numbers or bullets typed as text instead of a real list",
                    f"p{index}",
                )
        has_image = paragraph._p.find(f".//{qn('w:drawing')}") is not None
        if (
            paragraph.paragraph_format.alignment is not None
            and style_name == "Normal"
            and not has_image
        ):
            flag("direct-alignment", "Alignment set per paragraph, not in the style", f"p{index}")

    if len([font for font, chars in body_fonts.items() if chars > 40]) > 1:
        top = ", ".join(
            f"{name or '?'} {size or '?'}pt" for (name, size), _ in body_fonts.most_common(4)
        )
        flag("mixed-body-fonts", f"Body text mixes fonts or sizes: {top}", "body")

    for number, shape in enumerate(document.inline_shapes):
        doc_pr = shape._inline.docPr
        if not (doc_pr.get("descr") or "").strip():
            flag("image-alt-text", "Image without alternative text", f"image {number}")

    for number, table in enumerate(document.tables):
        first = table.rows[0] if table.rows else None
        if first is None:
            continue
        tr_pr = first._tr.trPr
        repeats = tr_pr is not None and tr_pr.find(qn("w:tblHeader")) is not None
        if not repeats and len(table.rows) > 8:
            flag(
                "table-header",
                "Long table whose header row does not repeat on each page",
                f"t{number}",
            )

    if document.element.body.find(f".//{qn('w:ins')}") is not None:
        flag("tracked-changes", "Document still has tracked changes", "body")
    return sorted(issues.values(), key=lambda item: -item["count"])
